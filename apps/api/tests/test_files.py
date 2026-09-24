"""Phase 3 A: file sandbox, text extraction, deliverables, evidence and the claim check (no network)."""

from __future__ import annotations

import asyncio
import json
import os
from pathlib import Path

import pytest

from atlas.core import paths
from atlas.core.events import EventBus
from atlas.core.models import Evidence
from atlas.core.registry import AgentRegistry
from atlas.core.store import WorldStore, fold
from atlas.live import AgentLoader, FakeLLM, LiveConfig, LiveEngine, ModelConfig, NodeContext, PriceTable
from atlas.live.evidence import file_mentions, unverified_claims
from atlas.live.files import FileAccessError, FileSandbox, FileTools, extract_text, render_deliverable
from atlas.live.llm import call_text, tool_use
from atlas.live.sdk import FakeClaudeSDK, call, say

MID = "msn_test0001"


@pytest.fixture(scope="module")
def registry() -> AgentRegistry:
    return AgentRegistry.load()


@pytest.fixture
def roots(tmp_path: Path, monkeypatch) -> dict[str, Path]:
    corp = tmp_path / "Documents"
    personal = corp / "Personal"  # another node's root nested inside corporate's
    outside = tmp_path / "outside"
    for d in (corp / "sub", personal, outside):
        d.mkdir(parents=True)
    (corp / "notes.txt").write_text("hello corporate")
    (corp / "sub" / "data.csv").write_text("a,b\n1,2\n")
    (personal / "diary.txt").write_text("PERSONAL")
    (outside / "secret.txt").write_text("SECRET")
    monkeypatch.setenv("ATLAS_FILE_ROOTS_CORPORATE", str(corp))
    monkeypatch.setenv("ATLAS_FILE_ROOTS_PERSONAL", str(personal))
    monkeypatch.delenv("ATLAS_FILE_MAX_CHARS", raising=False)
    return {"corp": corp, "personal": personal, "outside": outside, "tmp": tmp_path}


def corp_tools() -> FileTools:
    return FileTools(FileSandbox.for_mission("corporate", MID))


def denied(sb: FileSandbox, raw: str) -> str:
    with pytest.raises(FileAccessError) as exc:
        sb.resolve(raw)
    return str(exc.value)


# ---------------------------------------------------------------------------
# sandbox
# ---------------------------------------------------------------------------


def test_sandbox_allows_roots_relative_and_windows_style(roots):
    sb = FileSandbox.for_mission("corporate", MID)
    corp = roots["corp"]
    assert sb.resolve(str(corp / "notes.txt")) == (corp / "notes.txt").resolve()
    assert sb.resolve("notes.txt") == (corp / "notes.txt").resolve()  # relative to a root
    assert sb.resolve("sub\\data.csv") == (corp / "sub" / "data.csv").resolve()  # backslashes
    assert sb.resolve(str(corp / "sub" / "data.csv").replace("/", "\\")).name == "data.csv"
    assert sb.resolve(f'"{corp / "notes.txt"}"').name == "notes.txt"  # quoted
    if os.name != "nt":
        assert "outside the allowed folders" in denied(sb, "C:\\Users\\me\\Documents\\x.xlsx")
        assert "outside the allowed folders" in denied(sb, "C:/Users/me/x.xlsx")
    assert "does not exist" in denied(sb, "missing.txt")


def test_sandbox_blocks_traversal_absolute_and_symlinks(roots):
    sb = FileSandbox.for_mission("corporate", MID)
    corp, outside = roots["corp"], roots["outside"]
    assert "outside" in denied(sb, "../outside/secret.txt")
    assert "outside" in denied(sb, str(corp / "sub" / ".." / ".." / "outside" / "secret.txt"))
    assert "outside" in denied(sb, str(outside / "secret.txt"))
    assert "outside" in denied(sb, "/etc/passwd")
    try:
        os.symlink(outside / "secret.txt", corp / "link.txt")
        os.symlink(outside, corp / "linkdir")
    except OSError:
        pytest.skip("symlinks not available")
    assert "outside" in denied(sb, "link.txt")
    assert "outside" in denied(sb, "linkdir/secret.txt")
    listing = corp_tools().list_files(str(corp), recursive=True).text
    assert "link" not in listing and "secret" not in listing


def test_sandbox_deny_list(roots, monkeypatch):
    corp = roots["corp"]
    for rel in (".env", ".env.local", ".git/config", ".ssh/id_rsa", "cert.pem", "server.key", "id_rsa.pub",
                ".claude/settings.json"):
        p = corp / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text("x")
    sb = FileSandbox.for_mission("corporate", MID)
    for rel in (".env", ".env.local", ".git/config", ".ssh/id_rsa", "cert.pem", "server.key", "id_rsa.pub",
                ".claude/settings.json"):
        assert "denied" in denied(sb, rel), rel
    listing = corp_tools().list_files(str(corp), recursive=True).text
    for name in (".env", ".git", "cert.pem", "server.key", "id_rsa", ".ssh"):
        assert name not in listing
    assert "notes.txt" in listing and "sub/data.csv" in listing


def test_sandbox_atlas_local_rules(roots, monkeypatch):
    corp = roots["corp"]
    local = corp / "atlas-local"  # an ATLAS_LOCAL_DIR inside a read root
    monkeypatch.setenv("ATLAS_LOCAL_DIR", str(local))
    files = {
        "context/corporate/company.md": True,
        "context/personal/me.md": False,
        "atlas.db": False,
        "agents/mercato.md": False,
        f"missions/{MID}/attachments/brief.pdf": True,
        "missions/msn_other/attachments/theirs.pdf": False,
        f"outputs/corporate/{MID}/memo.md": True,
        "outputs/corporate/msn_other/memo.md": False,
        f"outputs/personal/{MID}/memo.md": False,
    }
    for rel in files:
        (local / rel).parent.mkdir(parents=True, exist_ok=True)
        (local / rel).write_text("x")
    sb = FileSandbox.for_mission("corporate", MID)
    for rel, ok in files.items():
        if ok:
            assert sb.resolve(str(local / rel)).name == Path(rel).name
        else:
            assert "denied" in denied(sb, str(local / rel)), rel
    assert sb.resolve("brief.pdf").name == "brief.pdf"  # a bare attachment name works
    listing = corp_tools().list_files(str(corp), recursive=True).text
    assert "atlas.db" not in listing and "me.md" not in listing and "company.md" in listing


def test_other_nodes_roots_are_unreachable(roots, monkeypatch):
    corp, personal = roots["corp"], roots["personal"]
    sb = FileSandbox.for_mission("corporate", MID)
    assert "another node" in denied(sb, str(personal / "diary.txt"))
    assert "another node" in denied(sb, "Personal/diary.txt")
    listing = corp_tools().list_files(str(corp), recursive=True).text
    assert "diary" not in listing and "Personal" not in listing
    assert corp_tools().search_files("diary").text.startswith("Search 'diary': 0 match")
    psb = FileSandbox.for_mission("personal", MID)
    assert psb.resolve("diary.txt").name == "diary.txt"
    assert "outside" in denied(psb, str(corp / "notes.txt"))
    # a node without configured roots reads only its mission folders
    monkeypatch.delenv("ATLAS_FILE_ROOTS_PERSONAL")
    assert FileSandbox.for_mission("personal", MID).user_roots == []


def test_size_cap_truncation_and_paging(roots, monkeypatch):
    corp = roots["corp"]
    big = corp / "big.txt"
    with open(big, "wb") as fh:
        fh.truncate(26 * 1024 * 1024)
    res = corp_tools().run("read_file", {"path": "big.txt"})
    assert not res.ok and "larger than 25 MB" in res.text and res.kind == "file_read"

    monkeypatch.setenv("ATLAS_FILE_MAX_CHARS", "1000")
    body = "".join(f"{i:04d}" for i in range(625))  # 2500 chars
    (corp / "long.md").write_text(body)
    tools = corp_tools()
    first = tools.read_file("long.md")
    assert "characters 0–1,000 of 2,500" in first.text and "offset=1000" in first.text
    assert first.text.endswith(body[:1000])
    second = tools.read_file("long.md", offset=1000)
    assert second.text.endswith(body[1000:2000]) and "offset=2000" in second.text
    last = tools.read_file("long.md", offset=2000, max_chars=99999)  # capped at ATLAS_FILE_MAX_CHARS
    assert last.text.endswith(body[2000:]) and "offset=" not in last.text
    small = tools.read_file("long.md", offset=10, max_chars=20)
    assert small.text.endswith(body[10:30]) and small.detail == "chars 10-30 of 2500"


def test_list_and_search_are_capped_and_skip_heavy_folders(roots):
    corp = roots["corp"]
    many = corp / "many"
    many.mkdir()
    for i in range(620):
        (many / f"f{i:03d}.txt").write_text("x")
    (corp / "node_modules" / "pkg").mkdir(parents=True)
    (corp / "node_modules" / "pkg" / "rocks.txt").write_text("x")
    (corp / "Rocks_Q3.xlsx").write_bytes(b"x")
    tools = corp_tools()
    res = tools.list_files(str(many))
    assert "500 entries (truncated at 500)" in res.text and res.detail == "500 entries"
    assert tools.list_files(str(corp), pattern="*.csv", recursive=True).text.count("\n") == 1
    hits = tools.search_files("rocks")
    assert "Rocks_Q3.xlsx" in hits.text and "node_modules" not in hits.text
    assert len(tools.search_files("f*.txt").text.splitlines()) == 101  # header + 100 hits
    roots_listing = tools.list_files()
    assert str(corp.resolve()) in roots_listing.text and "Mission attachments" in roots_listing.text


# ---------------------------------------------------------------------------
# extraction
# ---------------------------------------------------------------------------


def test_extract_pdf(tmp_path):
    canvas = pytest.importorskip("reportlab.pdfgen.canvas")
    path = tmp_path / "brief.pdf"
    c = canvas.Canvas(str(path))
    c.drawString(72, 720, "Quarterly rocks: close Polanco deal")
    c.showPage()
    c.drawString(72, 720, "Second page text")
    c.save()
    text, kind = extract_text(path)
    assert "--- Page 1 ---" in text and "Polanco deal" in text and "Second page text" in text
    assert kind == "pdf, 2 pages"


def test_extract_xlsx_sheets_and_formulas(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    path = tmp_path / "Rocks_Q3.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rocks"
    ws.append(["Rock", "Owner", "Pct"])
    ws.append(["Close deal", "Ana", 0.5])
    ws.append(["Hire CFO", "Luis", 1])
    ws["C4"] = "=SUM(C2:C3)"  # no cached value (openpyxl never computes) -> formula text
    ws2 = wb.create_sheet("Budget")
    ws2.append(["Item", "MXN"])
    ws2.append(["Rent", 18000])
    wb.save(path)
    text, kind = extract_text(path)
    assert "## Sheet: Rocks" in text and "## Sheet: Budget" in text
    assert "Close deal | Ana | 0.5" in text and "Hire CFO | Luis | 1" in text
    assert "=SUM(C2:C3)" in text and "Rent | 18000" in text
    assert kind == "workbook, sheets: Rocks, Budget"
    only, _ = extract_text(path, "budget")
    assert "Rent" in only and "Close deal" not in only
    with pytest.raises(FileAccessError, match="no sheet 'Nope'.*Rocks, Budget"):
        extract_text(path, "Nope")

    xlsxwriter = pytest.importorskip("xlsxwriter")
    cached = tmp_path / "cached.xlsx"
    book = xlsxwriter.Workbook(str(cached))
    sheet = book.add_worksheet("Totals")
    sheet.write_row(0, 0, ["a", 10])
    sheet.write_row(1, 0, ["b", 20])
    sheet.write_formula(2, 1, "=SUM(B1:B2)", None, 30)  # cached value -> shown instead of the formula
    book.close()
    text, _ = extract_text(cached)
    assert "30" in text.splitlines()[-1] and "SUM" not in text


def test_extract_docx_pptx_csv_md(tmp_path):
    docx = pytest.importorskip("docx")
    doc = docx.Document()
    doc.add_heading("Memo", level=1)
    doc.add_paragraph("First paragraph.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Rock", "Owner"
    table.cell(1, 0).text, table.cell(1, 1).text = "Close deal", "Ana"
    doc.add_paragraph("After the table.")
    doc.save(tmp_path / "memo.docx")
    text, _ = extract_text(tmp_path / "memo.docx")
    assert text.index("# Memo") < text.index("First paragraph.") < text.index("| Rock | Owner |")
    assert "| Close deal | Ana |" in text and text.index("| Close deal") < text.index("After the table.")

    pptx = pytest.importorskip("pptx")
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Q3 Rocks"
    slide.placeholders[1].text = "Close the Polanco deal"
    prs.save(tmp_path / "deck.pptx")
    text, kind = extract_text(tmp_path / "deck.pptx")
    assert "--- Slide 1 ---" in text and "Q3 Rocks" in text and "Polanco" in text and kind.endswith("1 slides")

    (tmp_path / "t.csv").write_text("name,rent\nPolanco,18000\n")
    (tmp_path / "n.md").write_text("# Título\nñandú")
    (tmp_path / "w.txt").write_bytes("café".encode("cp1252"))
    (tmp_path / "bin.dat").write_bytes(b"\x00\x01\x02")
    assert extract_text(tmp_path / "t.csv")[0] == "name,rent\nPolanco,18000\n"
    assert extract_text(tmp_path / "n.md")[0] == "# Título\nñandú"
    assert extract_text(tmp_path / "w.txt")[0] == "café"
    with pytest.raises(FileAccessError, match="binary"):
        extract_text(tmp_path / "bin.dat")
    (tmp_path / "old.xls").write_bytes(b"x")
    with pytest.raises(FileAccessError, match="not supported"):
        extract_text(tmp_path / "old.xls")


# ---------------------------------------------------------------------------
# deliverables
# ---------------------------------------------------------------------------


def test_write_deliverable_formats_and_never_overwrites(roots):
    tools = corp_tools()
    out = paths.outputs_dir("corporate", MID)
    cases = [
        ({"filename": "memo", "format": "md", "content": "# Hi"}, "memo.md", "markdown"),
        ({"filename": "memo.md", "format": "md", "content": "# Again"}, "memo (2).md", "markdown"),
        ({"filename": "notes", "format": "txt", "content": "plain"}, "notes.txt", "text"),
        ({"filename": "t", "format": "csv", "sheets": {"S": [["a", "b"], [1, 2]]}}, "t.csv", "file"),
        ({"filename": "d", "format": "json", "content": '{"a": 1}'}, "d.json", "json"),
        ({"filename": "resumen", "format": "xlsx", "sheets": {"Rocks": [["Rock", "Pct"], ["A", 0.5]],
                                                               "Otra": [["x"]]}}, "resumen.xlsx", "file"),
        ({"filename": "../../evil", "format": "docx",
          "content": "# Informe\n\nTexto **clave**.\n\n- uno\n- dos\n\n| A | B |\n|---|---|\n| 1 | 2 |"},
         "evil.docx", "file"),
    ]
    for args, name, kind in cases:
        res = tools.run("write_deliverable", args)
        assert res.ok, res.text
        att = res.attachment
        assert att is not None and att.name == name and att.kind == kind
        assert att.download_url == f"/missions/{MID}/files/outputs/{name.replace(' ', '%20').replace('(', '%28').replace(')', '%29')}"
        assert (out / name).is_file() and att.size_bytes == (out / name).stat().st_size
        assert res.kind == "file_written" and Path(res.ref) == (out / name).resolve()
    assert (out / "memo.md").read_text() == "# Hi"  # the first file was not overwritten
    assert (out / "t.csv").read_text() == "a,b\n1,2\n"
    assert json.loads((out / "d.json").read_text()) == {"a": 1}
    xlsx, _ = extract_text(out / "resumen.xlsx")
    assert "## Sheet: Rocks" in xlsx and "A | 0.5" in xlsx and "## Sheet: Otra" in xlsx
    docx_text, _ = extract_text(out / "evil.docx")
    assert "# Informe" in docx_text and "Texto clave." in docx_text and "| 1 | 2 |" in docx_text
    assert "- uno" in docx_text
    assert not list(roots["corp"].glob("**/evil.docx"))  # outputs only, never next to the user's files

    bad = tools.run("write_deliverable", {"filename": "x", "format": "json", "content": "{nope"})
    assert not bad.ok and "not valid JSON" in bad.text and bad.kind == "file_written"
    assert not tools.run("write_deliverable", {"filename": "x", "format": "exe", "content": "x"}).ok
    with pytest.raises(FileAccessError):
        render_deliverable("xlsx", None, None)


# ---------------------------------------------------------------------------
# claim check
# ---------------------------------------------------------------------------


def test_file_mentions_and_claim_check():
    assert file_mentions("read C:\\Docs\\Rocks Q3.xlsx") == ["Q3.xlsx"]  # names with spaces: last word
    mentions = file_mentions("Exported resumen.xlsx, read /data/Rocks_Q3.xlsx; see https://x.com/a.pdf")
    assert mentions == ["resumen.xlsx", "/data/Rocks_Q3.xlsx"]
    ev = [Evidence(mission_id="m", task_id="t1", agent_id="sofia", kind="file_read", ref="/home/u/Rocks_Q3.xlsx")]
    other = [Evidence(mission_id="m", task_id="t0", agent_id="argos", kind="file_written",
                      ref="/out/memo (2).docx", detail="3 KB, requested as memo.docx")]
    assert unverified_claims(["Read Rocks_Q3.xlsx", "Exported resumen.xlsx"], [], task_evidence=ev,
                             mission_evidence=ev) == ["resumen.xlsx"]
    # inputs may come from other tasks of the mission; actions must be this task's own
    assert unverified_claims([], ["memo.docx from the upstream task"], task_evidence=ev,
                             mission_evidence=ev + other) == []
    assert unverified_claims(["Wrote memo.docx"], [], task_evidence=ev, mission_evidence=ev + other) == ["memo.docx"]
    # node context files the system provided count as backed; failed actions don't
    assert unverified_claims([], ["company.md"], task_evidence=[], mission_evidence=[],
                             provided="### company.md\nmargins") == []
    failed = [ev[0].model_copy(update={"ok": False})]
    assert unverified_claims(["Read Rocks_Q3.xlsx"], [], task_evidence=failed, mission_evidence=failed)


# ---------------------------------------------------------------------------
# live missions: evidence + deliverables on the report, claim check (both backends)
# ---------------------------------------------------------------------------

MODELS = ModelConfig(orchestrator="o", default="s", fast="f")


def _plan_tasks() -> list[dict]:
    return [
        {"ref": "rocks", "title": "Review rocks", "description": "Read the Q3 rocks workbook, write a memo",
         "assigned_to": "oracle"},
        {"ref": "sum", "title": "Summarize", "description": "Summarize", "assigned_to": "alfred"},
    ]


def _workbook(root: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Rocks"
    wb.active.append(["Rock", "Owner"])
    wb.active.append(["Close Polanco", "Ana"])
    wb.save(root / "Rocks_Q3.xlsx")


FILE_STEPS = [
    ("list_files", {"path": "{root}"}),
    ("read_file", {"path": "Rocks_Q3.xlsx"}),
    ("write_deliverable", {"filename": "memo", "format": "docx", "content": "# Memo\n\n- Close Polanco (Ana)"}),
]
GOOD_REPORT = {"asked_to": "rocks", "actions_taken": ["Read Rocks_Q3.xlsx", "Wrote memo.docx"],
               "inputs_used": ["Rocks_Q3.xlsx"],
               "findings": [{"kind": "FACT", "statement": "Ana owns Close Polanco", "confidence": "HIGH"}],
               "confidence": "HIGH"}
LYING_REPORT = {"asked_to": "sum", "actions_taken": ["Exported resumen.xlsx"],
                "findings": [{"kind": "FACT", "statement": "Summary exported", "confidence": "HIGH"}],
                "confidence": "HIGH"}
MISSION_REPORT = {"executive_summary": "Done.", "objective_status": "ACHIEVED",
                  "key_findings": [{"kind": "FACT", "statement": "ok", "confidence": "HIGH"}]}


def _check_mission(store: WorldStore, initial, root: Path) -> None:
    snap = store.snapshot()
    reports = {r.agent_id: r for r in snap.agent_reports}
    good, bad = reports["oracle"], reports["alfred"]
    assert [e.kind for e in good.evidence] == ["file_listed", "file_read", "file_written"]
    assert all(e.ok for e in good.evidence) and good.evidence[1].ref == str((root / "Rocks_Q3.xlsx").resolve())
    assert good.evidence[1].task_id == good.task_id
    assert good.confidence == "HIGH" and not any(x.startswith("Unverified") for x in good.limitations)
    assert [d.name for d in good.deliverables] == ["memo.docx"]
    mid = snap.missions[0].id
    assert good.deliverables[0].download_url == f"/missions/{mid}/files/outputs/memo.docx"
    assert (paths.outputs_dir("corporate", mid) / "memo.docx").is_file()
    assert bad.evidence == [] and bad.deliverables == []
    assert bad.confidence == "LOW" and "Unverified: resumen.xlsx (no system record)" in bad.limitations
    events = store.bus.history()
    activities = [e.payload["state"].get("activity") for e in events
                  if e.type == "agent.state_changed" and e.agent_id == "oracle"]
    assert "Reading Rocks_Q3.xlsx" in activities and "Writing memo.docx" in activities
    feed = [e.summary for e in events if e.type == "evidence.recorded"]
    assert "ORACLE read Rocks_Q3.xlsx" in feed and "ORACLE wrote memo.docx" in feed
    assert len(snap.evidence) == 3
    final = snap.mission_reports[-1]
    assert [d.name for d in final.deliverables] == ["memo.docx"]
    assert fold(initial, events).model_dump(mode="json") == snap.model_dump(mode="json")


def test_live_mission_files_evidence_api(registry, roots):
    root = roots["corp"]
    _workbook(root)
    llm = FakeLLM()
    llm.when(lambda kw: "STAGE: PLANNING" in call_text(kw), tool_use("create_plan", {"tasks": _plan_tasks()}))
    llm.when(lambda kw: "Your task: Review rocks" in call_text(kw),
             *[tool_use(n, {k: v.format(root=root) for k, v in a.items()}) for n, a in FILE_STEPS],
             tool_use("submit_report", GOOD_REPORT))
    llm.when(lambda kw: "Your task: Summarize" in call_text(kw), tool_use("submit_report", LYING_REPORT))
    llm.when(lambda kw: "STAGE: REVIEW" in call_text(kw), tool_use("request_followups", {"tasks": []}))
    llm.when(lambda kw: "STAGE: CONSOLIDATION" in call_text(kw), tool_use("submit_mission_report", MISSION_REPORT))

    async def go():
        store = WorldStore(registry, EventBus())
        engine = LiveEngine(store, loader=AgentLoader(registry, MODELS),
                            context=NodeContext(registry, local_dir=Path("/nonexistent")),
                            config=LiveConfig(models=MODELS, max_turns=6), llm=llm, prices=PriceTable())
        initial = store.snapshot()
        m = await engine.start("Revisar rocks Q3", "corporate")
        await asyncio.wait_for(engine.wait(m.id), 5)
        return store, initial

    store, initial = asyncio.run(go())
    _check_mission(store, initial, root)
    rocks_calls = [c for c in llm.calls if "Your task: Review rocks" in call_text(c)]
    first = call_text(rocks_calls[0])
    assert str(root.resolve()) in first and "write_deliverable" in first  # files note + protocol
    assert {t["name"] for t in rocks_calls[0]["tools"]} >= {"list_files", "search_files", "read_file",
                                                            "write_deliverable"}
    read_result = json.dumps(rocks_calls[-1]["messages"], ensure_ascii=False)  # the list is shared
    assert "Close Polanco | Ana" in read_result


def test_live_mission_files_evidence_subscription(registry, roots):
    root = roots["corp"]
    _workbook(root)
    sdk = FakeClaudeSDK()
    sdk.when(lambda s: "STAGE: PLANNING" in s.prompt, [call("create_plan", {"tasks": _plan_tasks()}), say("ok")])
    sdk.when(lambda s: "Your task: Review rocks" in s.prompt, [
        *[call(n, {k: v.format(root=root) for k, v in a.items()}) for n, a in FILE_STEPS],
        call("submit_report", GOOD_REPORT), say("done")])
    sdk.when(lambda s: "Your task: Summarize" in s.prompt, [call("submit_report", LYING_REPORT), say("done")])
    sdk.when(lambda s: "STAGE: REVIEW" in s.prompt, [call("request_followups", {"tasks": []}), say("ok")])
    sdk.when(lambda s: "STAGE: CONSOLIDATION" in s.prompt, [call("submit_mission_report", MISSION_REPORT),
                                                           say("ok")])

    async def go():
        store = WorldStore(registry, EventBus())
        engine = LiveEngine(store, loader=AgentLoader(registry, MODELS),
                            context=NodeContext(registry, local_dir=Path("/nonexistent")),
                            config=LiveConfig(models=MODELS, max_turns=8), prices=PriceTable(),
                            backend="subscription", sdk_query=sdk)
        initial = store.snapshot()
        m = await engine.start("Revisar rocks Q3", "corporate")
        await asyncio.wait_for(engine.wait(m.id), 5)
        return store, initial

    store, initial = asyncio.run(go())
    _check_mission(store, initial, root)
    sess = next(s for s in sdk.sessions if "Your task: Review rocks" in s.prompt)
    assert [r[0] for r in sess.results][:3] == ["list_files", "read_file", "write_deliverable"]
    assert not any(r[2] for r in sess.results)
    assert "Close Polanco | Ana" in sess.results[1][1]
    # our file tools are MCP tools; Claude Code's own Read/Glob/Grep stay disabled
    assert "mcp__atlas__read_file" in sess.options.allowed_tools
    assert {"Read", "Glob", "Grep", "Write"} <= set(sess.options.disallowed_tools)
    assert sess.options.tools == [] and "mcp__atlas__list_files" in sess.system


class _Block:
    """A server-tool content block (has model_dump, like the SDK's)."""

    def __init__(self, **kw):
        self.__dict__.update(kw)

    def model_dump(self, **_):
        return dict(self.__dict__)


def test_evidence_for_consult_approval_and_web_api(registry, roots):
    from atlas.live.llm import FakeBlock, FakeMessage, text

    web = FakeMessage([
        _Block(type="server_tool_use", id="srv_1", name="web_search", input={"query": "Polanco rents"}),
        _Block(type="web_search_tool_result", tool_use_id="srv_1", content=[]),
        _Block(type="server_tool_use", id="srv_2", name="web_fetch", input={"url": "https://x.test/a"}),
        _Block(type="web_fetch_tool_result", tool_use_id="srv_2",
               content={"type": "web_fetch_tool_result_error", "error_code": "url_not_accessible"}),
        FakeBlock("tool_use", id="toolu_c", name="consult", input={"agent_id": "oracle", "question": "Rate?"}),
    ], "tool_use")
    llm = FakeLLM()
    llm.when(lambda kw: "(consultation)" in call_text(kw), text("10%"))
    llm.when(lambda kw: "STAGE: PLANNING" in call_text(kw), tool_use("create_plan", {"tasks": [
        {"ref": "a", "title": "Research", "description": "r", "assigned_to": "sofia", "requires_approval": True},
        {"ref": "b", "title": "Check", "description": "c", "assigned_to": "argos"}]}))
    llm.when(lambda kw: "Your task: Research" in call_text(kw), web,
             tool_use("request_approval", {"reason": "INSUFFICIENT_INFORMATION", "title": "Need data",
                                           "detail": "ok?"}),
             tool_use("submit_report", {"actions_taken": ["searched"], "findings": [], "confidence": "MEDIUM"}))
    llm.when(lambda kw: "Your task: Check" in call_text(kw),
             tool_use("submit_report", {"actions_taken": ["checked"], "findings": [], "confidence": "MEDIUM"}))
    llm.when(lambda kw: "STAGE: REVIEW" in call_text(kw), tool_use("request_followups", {"tasks": []}))
    llm.when(lambda kw: "STAGE: CONSOLIDATION" in call_text(kw), tool_use("submit_mission_report", MISSION_REPORT))

    async def go():
        store = WorldStore(registry, EventBus())
        engine = LiveEngine(store, loader=AgentLoader(registry, MODELS),
                            context=NodeContext(registry, local_dir=Path("/nonexistent")),
                            config=LiveConfig(models=MODELS, max_turns=6, web_search=True), llm=llm,
                            prices=PriceTable())
        m = await engine.start("Research", "corporate")
        deadline = asyncio.get_running_loop().time() + 3
        while not any(a.state == "PENDING" for a in store.snapshot().approvals):
            assert asyncio.get_running_loop().time() < deadline
            await asyncio.sleep(0.001)
        await store.decide_approval(store.snapshot().approvals[0].id, "REJECTED", note="no")
        await asyncio.wait_for(engine.wait(m.id), 5)
        return store

    store = asyncio.run(go())
    report = next(r for r in store.snapshot().agent_reports if r.agent_id == "sofia")
    got = [(e.kind, e.ref, e.ok) for e in report.evidence]
    apr = store.snapshot().approvals[0]
    assert got == [("web_search", "Polanco rents", True), ("web_fetch", "https://x.test/a", False),
                   ("consult", "oracle", True), ("approval", apr.id, True)]
    assert report.evidence[1].detail == "url_not_accessible"
    assert report.evidence[3].detail.startswith("REJECTED: Need data")
    feed = [e.summary for e in store.bus.history() if e.type == "evidence.recorded"]
    assert "SOFIA searched the web · Polanco rents" in feed and "SOFIA consulted ORACLE" in feed
