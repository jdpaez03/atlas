"""SCRIBE and institutional documents (docs/PUBLISHING.md): brand kit, spec validation, the PDF and deck
renderers, and the publishing step at the end of a live mission (both backends)."""

from __future__ import annotations

import asyncio
import copy
import io
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pypdf import PdfReader

from atlas.core import paths
from atlas.core.events import EventBus
from atlas.core.registry import AgentRegistry
from atlas.core.store import WorldStore
from atlas.live import AgentLoader, FakeLLM, LiveConfig, LiveEngine, ModelConfig, NodeContext, PriceTable
from atlas.live.llm import call_text, tool_use
from atlas.live.sdk import FakeClaudeSDK, call, say
from atlas.publish.brand import NEUTRAL_COLORS, brand_dir, deck_base_from_template, load_brand
from atlas.publish.deck import render_deck
from atlas.publish.pdf import DocMeta, render_pdf
from atlas.publish.spec import accent_runs, normalize, spec_text

SPEC = {
    "doc_kind": "Estudio de mercado",
    "title": "Polanco *2BR*: precio y demanda",
    "subtitle": "Comité de inversión",
    "summary": ["El precio de lista sostenible es 95,000 MXN/m² con absorción de 1.8 unidades por mes.",
                "La oferta directa suma 4 proyectos."],
    "highlights": [{"label": "Precio", "value": "95,000", "note": "MXN/m²"},
                   {"label": "Absorción", "value": "1.8", "note": "unidades/mes"}],
    "sections": [
        {"title": "Oferta *competidora*", "blocks": [
            {"type": "paragraph", "text": "Cuatro proyectos compiten en un radio de 2 km."},
            {"type": "table", "table": {"columns": ["Proyecto", "Precio MXN/m²", "Unidades"],
                                        "rows": [["Torre A", 95000, 40], ["Torre B", 88000, 36],
                                                 ["Torre C", 101000, 22, "extra cell dropped"], ["Torre D"]],
                                        "source": "Reporte ORACLE"}},
            {"type": "chart", "chart": {"chart_type": "bar", "categories": ["Torre A", "Torre B"],
                                        "series": [{"name": "MXN/m²", "values": [95000, 88000]}]}},
        ]},
        {"title": "Riesgos", "blocks": [
            {"type": "bullets", "items": ["Tasa hipotecaria.", "Inventario de 99,999 unidades en la zona."]},
            {"type": "callout", "title": "Supuesto", "text": "Lanzamiento en T1 2027."},
            {"type": "kpis", "kpis": [{"label": "Proyectos", "value": "4"}]},
        ]},
    ],
    "slides": [
        {"type": "statement", "text": "Polanco 2BR es viable a *95,000 MXN/m²*."},
        {"type": "kpis", "title": "Cifras *clave*", "kpis": [{"label": "Precio", "value": "95,000"}]},
        {"type": "section", "title": "Oferta"},
        {"type": "table", "title": "Comparables", "table": {"columns": ["Proyecto", "Precio"],
                                                             "rows": [["Torre A", 95000], ["Torre B", 88000]]}},
        {"type": "chart", "title": "Ventas", "chart": {"chart_type": "line", "categories": ["T1", "T2"],
                                                        "series": [{"name": "Base", "values": [5, 9]}]}},
        {"type": "bullets", "title": "Riesgos", "items": ["Tasa hipotecaria."], "notes": "Mencionar tasa."},
        {"type": "quote", "text": "La oferta es limitada.", "attribution": "ORACLE"},
    ],
    "sources": ["Reporte ORACLE"],
}


# ---------------------------------------------------------------------------
# brand kit
# ---------------------------------------------------------------------------


def _kit(node: str = "corporate") -> Path:
    from PIL import Image

    folder = brand_dir(node)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "brand.yaml").write_text(
        "company: Acme Desarrollos\nside_label: ACME\nfooter: Confidencial Acme\n"
        "colors: {primary: '#002A53', light: 'F5F5F5', ink: nope}\n", encoding="utf-8")
    img = Image.new("RGBA", (300, 200), (255, 255, 255, 0))
    img.paste((255, 255, 255, 255), (50, 50, 250, 150))
    img.save(folder / "logo_light.png")
    img.save(folder / "logo_dark.png")
    return folder


def test_neutral_brand_without_a_kit():
    b = load_brand("personal")
    assert not b.configured and b.folder is None and b.colors == NEUTRAL_COLORS
    assert b.logo_light is None and b.deck_base is None
    assert set(b.pdf_fonts) == {"light", "regular", "semibold", "italic"}


def test_brand_kit_is_loaded_and_sanitized():
    folder = _kit()
    b = load_brand("corporate")
    assert b.company == "Acme Desarrollos" and b.label == "ACME" and b.footer == "Confidencial Acme"
    assert b.c("primary") == "#002A53" and b.c("light") == "#F5F5F5"
    assert b.c("ink") == NEUTRAL_COLORS["ink"]  # invalid color → default
    assert b.logo_light == folder / "logo_light.png" and b.deck_base is None


def test_deck_base_from_a_macro_template(tmp_path):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0])
    buf = io.BytesIO()
    prs.save(buf)
    # make it a .potm: template content type + a vbaProject part referenced by the presentation
    src = tmp_path / "tpl.potm"
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as zin, zipfile.ZipFile(src, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
                    b"application/vnd.ms-powerpoint.template.macroEnabled.main+xml").replace(
                    b"</Types>", b'<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/>'
                                 b"</Types>")
            if item.filename == "ppt/_rels/presentation.xml.rels":
                data = data.replace(b"</Relationships>", b'<Relationship Id="rId99" Type="http://schemas.'
                                    b'microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/>'
                                    b"</Relationships>")
            zout.writestr(item, data)
        zout.writestr("ppt/vbaProject.bin", b"macro")
    out = deck_base_from_template(src, tmp_path / "deck_base.pptx")
    base = Presentation(str(out))
    assert len(base.slides) == 0 and len(base.slide_layouts) > 0
    names = zipfile.ZipFile(out).namelist()
    assert not any("vba" in n.lower() for n in names)


# ---------------------------------------------------------------------------
# spec
# ---------------------------------------------------------------------------


def test_normalize_cleans_and_reports_errors():
    spec, errors = normalize(copy.deepcopy(SPEC))
    assert errors == []
    rows = spec["sections"][0]["blocks"][1]["table"]["rows"]
    assert rows[2] == ["Torre C", 101000, 22] and rows[3] == ["Torre D", "", ""]  # squared to the columns
    _, errs = normalize({"title": "", "summary": [], "sections": [{"title": "x", "blocks": [
        {"type": "chart", "chart": {"chart_type": "bar", "categories": ["a", "b"],
                                    "series": [{"name": "s", "values": [1]}]}},
        {"type": "table", "table": {"columns": ["a"], "rows": []}}]}], "slides": []})
    joined = " | ".join(errs)
    for needle in ("title is required", "summary is required", "has 1 values for 2 categories",
                   "table has no rows", "at least one section", "the deck needs slides"):
        assert needle in joined, needle
    assert normalize(None)[1]
    assert accent_runs("Nuestra *huella* hoy") == [("Nuestra ", False), ("huella", True), (" hoy", False)]
    assert "99,999" in spec_text(spec) and "88000" in spec_text(spec)


# ---------------------------------------------------------------------------
# renderers
# ---------------------------------------------------------------------------


def test_render_pdf_and_deck(tmp_path):
    _kit()
    brand = load_brand("corporate")
    spec, _ = normalize(copy.deepcopy(SPEC))
    pdf = render_pdf(spec, brand, DocMeta("24 de septiembre de 2026", {"Verificación": ["AUDITOR: todo PASS"]}),
                     tmp_path / "r.pdf")
    reader = PdfReader(str(pdf))
    text = "\n".join(p.extract_text() for p in reader.pages)
    assert len(reader.pages) >= 4  # cover, summary + sections, notes
    for needle in ("Polanco", "Resumen", "Oferta", "Torre C", "101,000", "Notas", "AUDITOR: todo PASS",
                   "Confidencial Acme"):
        assert needle in text, needle
    assert reader.metadata.title == "Polanco 2BR: precio y demanda"

    deck = render_deck(spec, brand, "24 de septiembre de 2026", ["Acme Desarrollos"], tmp_path / "d.pptx")
    prs = Presentation(str(deck))
    assert len(prs.slides) == len(spec["slides"]) + 2  # cover + body + closing
    texts = ["\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
    assert "Polanco 2BR: precio y demanda" in texts[0].replace("\n", " ")
    table = next(sh for sh in prs.slides[4].shapes if sh.has_table).table
    assert table.cell(0, 0).text == "Proyecto" and table.cell(1, 1).text == "95,000"
    assert any(sh.has_chart for sh in prs.slides[5].shapes)
    assert prs.slides[6].notes_slide.notes_text_frame.text == "Mencionar tasa."
    assert any(sh.shape_type == 13 for sh in prs.slides[0].shapes)  # the logo picture on the cover


def test_render_without_a_kit(tmp_path):
    spec, _ = normalize(copy.deepcopy(SPEC))
    brand = load_brand("personal")
    assert render_pdf(spec, brand, DocMeta("1 de enero de 2027"), tmp_path / "n.pdf").stat().st_size > 1000
    assert len(Presentation(str(render_deck(spec, brand, "x", [], tmp_path / "n.pptx"))).slides) == 9


# ---------------------------------------------------------------------------
# in a mission
# ---------------------------------------------------------------------------

MODELS = ModelConfig(orchestrator="o", default="s", fast="f")
PLAN = [{"ref": "a", "title": "Price study", "description": "Comparables", "assigned_to": "oracle"},
        {"ref": "b", "title": "Summarize", "description": "Summarize", "assigned_to": "alfred"}]
ORACLE = {"asked_to": "a", "findings": [
    {"kind": "FACT", "statement": "Torre A lists at 95,000 MXN/m²; Torre B 88,000; Torre C 101,000 with 22 units",
     "confidence": "HIGH"},
    {"kind": "FACT", "statement": "Absorption 1.8 units/month; 4 projects; Torre A 40 units, Torre B 36; "
                                  "sales 5 then 9 by T2; launch T1 2027", "confidence": "MEDIUM"}],
    "confidence": "HIGH"}
ALFRED = {"asked_to": "b", "findings": [{"kind": "FACT", "statement": "ok", "confidence": "HIGH"}],
          "confidence": "HIGH"}
FINAL = {"executive_summary": "Viable at 95,000 MXN/m².", "objective_status": "ACHIEVED",
         "key_findings": [{"kind": "FACT", "statement": "95,000 MXN/m²", "confidence": "HIGH"}]}


def _engine(store, reg, **kw):
    return LiveEngine(store, loader=AgentLoader(reg, MODELS), context=NodeContext(reg, local_dir=Path("/nonexistent")),
                      prices=PriceTable(), **kw)


def _script_api(llm: FakeLLM, publish_response=None) -> FakeLLM:
    llm.when(lambda kw: "STAGE: PLANNING" in call_text(kw), tool_use("create_plan", {"tasks": PLAN}))
    llm.when(lambda kw: "Your task: Price study" in call_text(kw), tool_use("submit_report", ORACLE))
    llm.when(lambda kw: "Your task: Summarize" in call_text(kw), tool_use("submit_report", ALFRED))
    llm.when(lambda kw: "STAGE: REVIEW" in call_text(kw), tool_use("request_followups", {"tasks": []}))
    llm.when(lambda kw: "STAGE: CONSOLIDATION" in call_text(kw), tool_use("submit_mission_report", FINAL))
    if publish_response is not None:
        llm.when(lambda kw: "STAGE: PUBLISHING" in call_text(kw), publish_response)
    return llm


def _run(engine, publish=True):
    async def go():
        m = await engine.start("Estudio Polanco 2BR", "corporate", publish=publish)
        await asyncio.wait_for(engine.wait(m.id), 30)
        return m.id

    return asyncio.run(go())


def test_scribe_publishes_documents_at_the_end_api():
    _kit()
    reg = AgentRegistry.load()
    store = WorldStore(reg, EventBus())
    llm = _script_api(FakeLLM(), tool_use("submit_documents", copy.deepcopy(SPEC)))
    mid = _run(_engine(store, reg, config=LiveConfig(models=MODELS, max_turns=4, audit=False), llm=llm))
    report = store.mission_reports_for(mid)[-1]
    names = [d.name for d in report.documents]
    assert names == ["ACME_Polanco_2BR_precio_y_demanda_Reporte.pdf",
                     "ACME_Polanco_2BR_precio_y_demanda_Presentacion.pptx"]
    folder = paths.outputs_dir("corporate", mid)
    assert all((folder / n).is_file() for n in names)
    assert report.documents[0].download_url == f"/missions/{mid}/files/outputs/{names[0]}"
    ev = [e for e in store.evidence_for(mid) if e.agent_id == "scribe"]
    assert [e.kind for e in ev] == ["file_written", "file_written"] and "PDF institucional" in ev[0].detail
    # the invented figure (99,999) is flagged in the log and in the PDF's verification notes
    log = [e.summary or "" for e in store.bus.history() if e.type == "log"]
    assert any("SCRIBE · 1 figure(s)" in line and "99999" in line for line in log)
    pdf_text = " ".join(" ".join(p.extract_text().split()) for p in PdfReader(str(folder / names[0])).pages)
    assert "Cifra sin respaldo rastreable en los reportes: 99999" in pdf_text
    call_ = next(c for c in llm.calls if "STAGE: PUBLISHING" in call_text(c))
    assert "Torre A lists at 95,000" in call_text(call_) and "Organization: Acme Desarrollos" in call_text(call_)
    board = [e.summary for e in store.bus.history() if e.type == "agent.state_changed" and e.agent_id == "scribe"]
    assert "SCRIBE · COMPLETED · Published 2 document(s) · 1 figure(s) to check" in board


def test_publish_off_per_mission_and_failures_never_lose_the_report():
    reg = AgentRegistry.load()
    store = WorldStore(reg, EventBus())
    llm = _script_api(FakeLLM())  # no PUBLISHING response
    cfg = LiveConfig(models=MODELS, max_turns=4, audit=False)
    off = _run(_engine(store, reg, config=cfg, llm=llm), publish=False)
    assert store.mission_reports_for(off)[-1].documents == []
    assert not any("STAGE: PUBLISHING" in call_text(c) for c in llm.calls)

    store2 = WorldStore(reg, EventBus())
    llm2 = _script_api(FakeLLM())
    failed = _run(_engine(store2, reg, config=cfg, llm=llm2))
    report = store2.mission_reports_for(failed)[-1]
    assert report.documents == [] and report.executive_summary == "Viable at 95,000 MXN/m²."
    assert any("Documents not produced" in (e.summary or "") for e in store2.bus.history())
    assert store2.mission(failed).phase == "CLOSED"

    store3 = WorldStore(reg, EventBus())
    llm3 = _script_api(FakeLLM())
    globally_off = _run(_engine(store3, reg, config=LiveConfig(models=MODELS, max_turns=4, audit=False,
                                                               publish=False), llm=llm3))
    assert store3.mission_reports_for(globally_off)[-1].documents == []
    assert not any("Documents not produced" in (e.summary or "") for e in store3.bus.history())


def test_scribe_publishes_on_the_subscription_backend():
    _kit()
    reg = AgentRegistry.load()
    store = WorldStore(reg, EventBus())
    sdk = FakeClaudeSDK()
    sdk.when(lambda s: "STAGE: PLANNING" in s.prompt, [call("create_plan", {"tasks": PLAN}), say("ok")])
    sdk.when(lambda s: "Your task: Price study" in s.prompt, [call("submit_report", ORACLE), say("done")])
    sdk.when(lambda s: "Your task: Summarize" in s.prompt, [call("submit_report", ALFRED), say("done")])
    sdk.when(lambda s: "STAGE: REVIEW" in s.prompt, [call("request_followups", {"tasks": []}), say("ok")])
    sdk.when(lambda s: "STAGE: CONSOLIDATION" in s.prompt, [call("submit_mission_report", FINAL), say("ok")])
    sdk.when(lambda s: "STAGE: PUBLISHING" in s.prompt, [call("submit_documents", copy.deepcopy(SPEC)), say("ok")])
    mid = _run(_engine(store, reg, config=LiveConfig(models=MODELS, max_turns=8, audit=False),
                       backend="subscription", sdk_query=sdk))
    report = store.mission_reports_for(mid)[-1]
    assert [d.name.rsplit("_", 1)[-1] for d in report.documents] == ["Reporte.pdf", "Presentacion.pptx"]
    assert "scribe" in store.mission(mid).usage_by_agent


@pytest.mark.parametrize("fmt", ["pdf", "pptx"])
def test_formats_can_be_limited(fmt):
    reg = AgentRegistry.load()
    store = WorldStore(reg, EventBus())
    llm = _script_api(FakeLLM(), tool_use("submit_documents", copy.deepcopy(SPEC)))
    mid = _run(_engine(store, reg, config=LiveConfig(models=MODELS, max_turns=4, audit=False,
                                                      publish_formats=(fmt,)), llm=llm))
    assert [d.name.rsplit(".", 1)[-1] for d in store.mission_reports_for(mid)[-1].documents] == [fmt]
