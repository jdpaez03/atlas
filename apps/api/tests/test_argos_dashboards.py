"""ARGOS dashboards check (docs/ARGOS.md § Dashboards check) + attachment support in the mail sources.

Dashboards are generated here (xlsx with openpyxl, pdf with reportlab): two ISO weeks (S38, S39), three projects.
    Amāra      pdf; the S39 file is byte-identical to S38                          → identical_report
    Altavista  pdf in S38 only (its S39 attachment fails to download)             → missing_report
    Nativa     xlsx; Cimentación moved 30-sep → 30-oct, "Permiso de ocupación" row gone (ARGOS findings), and
               "Escrituradas 22/51" while 23 rows say Escriturado                 → kpi_mismatch
The mail source is a fake with attachments, the LLM a FakeLLM (api backend) whose findings include quotes that
are not verbatim (dropped).
"""

from __future__ import annotations

import asyncio
import base64
import io
import json
import shutil
from datetime import UTC, datetime, timedelta
from email.message import EmailMessage
from pathlib import Path

import httpx
import pytest
import yaml

from atlas.argos import dashboards as dash
from atlas.argos.checks import CheckContext, CheckNotConfigured
from atlas.core import paths
from atlas.core.events import EventBus
from atlas.core.models import EmailDraft
from atlas.core.registry import AgentRegistry
from atlas.core.store import WorldStore
from atlas.inbox.sources.base import (
    AttachmentMeta,
    AttachmentSource,
    MailMessage,
    MailSourceError,
    SourceStatus,
)
from atlas.inbox.sources.folder import FolderSource
from atlas.inbox.sources.graph import GraphSource
from atlas.live import AgentLoader, FakeLLM, LiveConfig, ModelConfig, NodeContext, PriceTable
from atlas.live.executor import ApiExecutor
from atlas.live.llm import Meter, call_text, tool_use
from atlas.live.runtime import MissionScope

FIXTURES = Path(__file__).parent / "fixtures" / "argos"
MODELS = ModelConfig(orchestrator="claude-opus-test", default="claude-sonnet-test", fast="claude-haiku-test")
NOW = datetime(2026, 9, 24, 16, 0, tzinfo=UTC)  # Thursday of ISO week 39
W38_AT = datetime(2026, 9, 18, 15, 0, tzinfo=UTC)  # Friday of week 38
W39_AT = datetime(2026, 9, 23, 15, 0, tzinfo=UTC)
W38, W39 = "2026-W38", "2026-W39"


def run(coro):
    return asyncio.run(coro)


# ---------------------------------------------------------------------------
# Generated dashboards
# ---------------------------------------------------------------------------


def nativa_xlsx(week: int, *, escrituradas: str, n_escriturado: int, hitos: list[tuple[str, str, str]]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append([f"Dashboard semanal Nativa S{week}"])
    ws.append(["Proyecto", "Nativa"])
    ws.append(["Escrituradas", escrituradas])
    ws.append(["Avance de obra", "58%" if week == 38 else "64%"])
    units = wb.create_sheet("Unidades")
    units.append(["Unidad", "Estatus", "Cliente"])
    for i in range(51):
        status = "Escriturado" if i < n_escriturado else ("Vendido" if i < 40 else "Disponible")
        units.append([f"N-{101 + i}", status, f"Cliente {i + 1}" if status != "Disponible" else None])
    ms = wb.create_sheet("Hitos")
    ms.append(["Hito", "Fecha", "Estatus"])
    for row in hitos:
        ms.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def project_pdf(title: str, kpis: list[str], hitos: list[tuple[str, str, str]]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"])] + [Paragraph(k, styles["Normal"]) for k in kpis] + [Spacer(1, 12)]
    table = Table([["Hito", "Fecha", "Estatus"], *[list(h) for h in hitos]])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    story.append(table)
    SimpleDocTemplate(buf, pagesize=A4).build(story)
    return buf.getvalue()


HITOS_38 = [("Cimentación", "30-sep-2026", "En proceso"), ("Estructura", "15-nov-2026", "Pendiente"),
            ("Permiso de ocupación", "15-dic-2026", "Pendiente")]
HITOS_39 = [("Cimentación", "30-oct-2026", "En proceso"), ("Estructura", "15-nov-2026", "Pendiente")]


# ---------------------------------------------------------------------------
# Fake mail source with attachments
# ---------------------------------------------------------------------------


class AttSource:
    name = "fake"

    def __init__(self, *, connected: bool = True):
        self.connected = connected
        self.messages: dict[str, MailMessage] = {}
        self.files: dict[str, list[tuple[AttachmentMeta, bytes]]] = {}
        self.broken: set[tuple[str, str]] = set()
        self.downloads: list[tuple[str, str]] = []

    def add(self, mid: str, subject: str, at: datetime, files: list[tuple[str, bytes]], *, broken=()) -> None:
        self.messages[mid] = MailMessage(
            id=mid, internet_message_id=f"<{mid}@x>", conversation_id=mid, subject=subject,
            sender="Gerente <gerente@constructora.mx>", to=["juan@corp.mx"], cc=[], received_at=at,
            is_from_me=False, web_link=None, preview="", has_attachments=bool(files),
        )
        self.files[mid] = [(AttachmentMeta(id=f"a{i}", name=n, size=len(b)), b) for i, (n, b) in enumerate(files)]
        self.broken |= {(mid, f"a{i}") for i, (n, _) in enumerate(files) if n in broken}

    async def status(self) -> SourceStatus:
        return SourceStatus(name="fake", connected=self.connected, detail="" if self.connected else "Not connected.",
                            account="juan@corp.mx")

    async def list_messages(self, since: datetime, limit: int = 200) -> list[MailMessage]:
        out = [m for m in self.messages.values() if m.received_at >= since]
        return sorted(out, key=lambda m: m.received_at, reverse=True)[:limit]

    async def get_body(self, message_id: str) -> str:
        raise AssertionError("the dashboards check must not read email bodies")

    async def create_outlook_draft(self, draft: EmailDraft) -> str | None:
        return None

    async def list_attachments(self, message_id: str) -> list[AttachmentMeta]:
        return [meta for meta, _ in self.files[message_id]]

    async def download_attachment(self, message_id: str, attachment_id: str) -> bytes:
        self.downloads.append((message_id, attachment_id))
        if (message_id, attachment_id) in self.broken:
            raise MailSourceError("attachment download failed", hint="retry later")
        return next(b for meta, b in self.files[message_id] if meta.id == attachment_id)


def mailbox() -> AttSource:
    src = AttSource()
    amara = project_pdf("Dashboard Amara - Semana 38", ["Avance de obra: 45%", "Fecha de corte: 18-sep-2026"],
                        [("Torre A", "30-nov-2026", "En proceso")])
    src.add("m-am-38", "Reporte semanal Amāra S38", W38_AT, [("Dashboard Amara S38.pdf", amara)])
    src.add("m-am-39", "Reporte semanal Amāra S39", W39_AT, [("Dashboard Amara S39.pdf", amara)])  # same bytes
    src.add("m-al-38", "Altavista · reporte semanal", W38_AT, [(
        "Altavista S38.pdf", project_pdf("Altavista - Semana 38", ["Avance de obra: 30%"],
                                         [("Excavación", "10-oct-2026", "En proceso")]))])
    src.add("m-al-39", "Altavista · reporte semanal", W39_AT, [("Altavista S39.pdf", b"%PDF-1.4 ...")],
            broken={"Altavista S39.pdf"})
    src.add("m-na-38", "Dashboard Torre N", W38_AT, [(
        "Nativa S38.xlsx", nativa_xlsx(38, escrituradas="20/51", n_escriturado=20, hitos=HITOS_38))])
    src.add("m-na-39", "Dashboard Torre N", W39_AT, [
        ("Nativa S39.xlsx", nativa_xlsx(39, escrituradas="22/51", n_escriturado=23, hitos=HITOS_39)),
        ("Nativa anexo S39.xlsx", b"this is not a workbook"),  # unreadable → a note, not a crash
        ("foto.jpg", b"\xff\xd8\xff"),  # not a dashboard type
    ])
    src.add("m-fact", "Factura septiembre", W39_AT, [("factura.pdf", b"%PDF-1.4 invoice")])  # no match
    src.add("m-old", "Reporte semanal Nativa S30", NOW - timedelta(days=40), [("Nativa S30.xlsx", b"old")])
    return src


FINDINGS = [
    {"kind": "moved_date", "severity": "HIGH", "title": "Cimentación se movió del 30-sep al 30-oct",
     "detail": "La fecha de Cimentación se recorrió 30 días.",
     "evidence": [{"version": "previous", "quote": "Cimentación | 30-sep-2026 | En proceso"},
                  {"version": "current", "quote": "Cimentación  30-oct-2026"}]},  # separators/space tolerated
    {"kind": "removed_row", "severity": "MEDIUM", "title": "Desapareció el hito Permiso de ocupación",
     "detail": "Estaba programado para el 15-dic-2026.",
     "evidence": [{"version": "previous", "quote": "Permiso de ocupación | 15-dic-2026 | Pendiente"},
                  {"version": "current", "quote": "Permiso de ocupación"}]},  # not in current: this quote drops
    {"kind": "value_change", "severity": "MEDIUM", "title": "Escrituradas subió a 25",
     "evidence": [{"version": "current", "quote": "Escrituradas | 25/51"}]},  # invented → finding dropped
    {"kind": "other", "severity": "LOW", "title": "Quote from the wrong version",
     "evidence": [{"version": "previous", "quote": "Cimentación | 30-oct-2026"}]},  # wrong version → dropped
]


def argos_llm() -> FakeLLM:
    llm = FakeLLM()
    llm.when(lambda kw: "STAGE: ARGOS DASHBOARDS" in call_text(kw),
             tool_use("record_dashboard_findings", {"findings": FINDINGS}), repeat=True)
    return llm


@pytest.fixture(scope="module")
def registry() -> AgentRegistry:
    return AgentRegistry.load()


@pytest.fixture
def local() -> Path:
    root = paths.local_dir()
    (root / "inbox").mkdir(parents=True, exist_ok=True)
    shutil.copy(FIXTURES / "digest_rules.yaml", root / "inbox" / "digest_rules.yaml")
    return root


def watch() -> dict:
    return yaml.safe_load((FIXTURES / "watch.yaml").read_text(encoding="utf-8"))


async def make_ctx(registry: AgentRegistry, src, llm: FakeLLM | None, *, now: datetime = NOW) -> CheckContext:
    store = WorldStore(registry, EventBus())
    mission = await store.create_mission("ARGOS watch · dashboards", "corporate", mode="live")
    scope = None
    if llm is not None:
        loader = AgentLoader(registry, MODELS)
        scope = MissionScope(
            store=store, meter=Meter(llm, store, mission.id, PriceTable()), config=LiveConfig(models=MODELS),
            context=NodeContext(registry, local_dir=Path("/nonexistent-atlas-local")), mission_id=mission.id,
            objective=mission.objective, node="corporate", agents={"argos": loader.resolve("argos")},
            orchestrator=loader.resolve(registry.orchestrator.id),
        )
    return CheckContext(store=store, scope=scope, mail=src, config=watch(), now=now, mission_id=mission.id,
                        executor=ApiExecutor() if llm is not None else None)


def by_kind(alerts) -> dict[str, list]:
    out: dict[str, list] = {}
    for a in alerts:
        out.setdefault(a.kind, []).append(a)
    return out


# ---------------------------------------------------------------------------
# End to end
# ---------------------------------------------------------------------------


def test_dashboards_end_to_end(registry, local):
    src, llm = mailbox(), argos_llm()
    ctx = run(make_ctx(registry, src, llm))
    result = run(dash.DashboardsCheck().run(ctx))
    kinds = by_kind(result.alerts)
    assert set(kinds) == {"missing_report", "identical_report", "kpi_mismatch", "moved_date", "removed_row"}

    missing = kinds["missing_report"]
    assert [(a.project, a.severity, a.fingerprint) for a in missing] == [
        ("Altavista", "HIGH", "dashboards:altavista:missing-report:2026-w39")]
    assert missing[0].evidence and missing[0].evidence[0].version == "previous"
    assert missing[0].evidence[0].quote == "Avance de obra: 30%"

    identical = kinds["identical_report"]
    assert [a.project for a in identical] == ["Amāra"]
    assert identical[0].fingerprint == "dashboards:amara:identical-report:2026-w39"
    assert "byte-identical" in identical[0].detail

    kpi = kinds["kpi_mismatch"]
    assert [(a.project, a.fingerprint) for a in kpi] == [("Nativa", "dashboards:nativa:kpi-mismatch:escrituradas")]
    assert "22" in kpi[0].title and "23" in kpi[0].title
    assert kpi[0].evidence[0].quote == "Escrituradas | 22/51"
    assert kpi[0].evidence[1].quote.startswith("N-101 | Escriturado")

    moved = kinds["moved_date"][0]
    assert moved.project == "Nativa" and moved.severity == "HIGH"
    assert moved.fingerprint == "dashboards:nativa:moved-date:cimentacion-se-movio-del-30-sep-al-30-oct"
    assert [(e.version, e.source) for e in moved.evidence] == [("previous", "Nativa S38.xlsx"),
                                                               ("current", "Nativa S39.xlsx")]
    removed = kinds["removed_row"][0]
    assert [e.version for e in removed.evidence] == ["previous"]  # the non-verbatim current quote was dropped
    assert all(a.title != "Escrituradas subió a 25" and a.kind != "other" for a in result.alerts)

    # one ARGOS call: Nativa (Amāra is identical, Altavista is missing); the prompt carries the code's hints
    assert len(llm.calls) == 1
    prompt = call_text(llm.calls[0])
    assert "Project: Nativa" in prompt and "Cimentación | 30-sep-2026" in prompt
    assert "date changed for 'Cimentación · Fecha': 2026-09-30 → 2026-10-30" in prompt
    assert "row no longer present: Permiso de ocupación | 15-dic-2026 | Pendiente" in prompt
    assert "Escrituradas' says 22 but the table lists 23" in prompt
    assert llm.calls[0]["tool_choice"] == {"type": "tool", "name": "record_dashboard_findings"}

    notes = "\n".join(result.notes)
    assert "Could not download 'Altavista S39.pdf' (Altavista)" in notes
    assert "Could not read 'Nativa anexo S39.xlsx'" in notes
    assert "quote(s) not found verbatim were dropped" in notes
    assert result.ok

    # files, index and snapshots on disk
    argos = local / "argos"
    assert (argos / "dashboards" / "Am_ra" / W38 / "Dashboard Amara S38.pdf").is_file()
    assert (argos / "dashboards" / "Nativa" / W39 / "Nativa S39.xlsx").is_file()
    assert not list((argos / "dashboards").rglob("factura.pdf")) and not list(argos.rglob("foto.jpg"))
    assert not list(argos.rglob("Nativa S30.xlsx"))  # older than lookback_days
    state = json.loads((argos / "state.json").read_text(encoding="utf-8"))["dashboards"]
    entry = state["attachments"]["m-na-39:a0"]
    assert entry["project"] == "Nativa" and entry["week"] == W39 and len(entry["sha256"]) == 64
    snap = json.loads((argos / "snapshots" / "Nativa" / f"{W39}.json").read_text(encoding="utf-8"))
    assert set(snap) >= {"kpis", "rows", "dates", "text_hash"}
    assert snap["kpis"]["Escrituradas"] == "22/51"
    assert {"label": "Cimentación · Fecha", "date": "2026-10-30", "raw": "30-oct-2026"} in snap["dates"]
    assert sum(1 for r in snap["rows"] if r["cells"].get("Estatus") == "Escriturado") == 23
    assert len([r for r in snap["rows"] if r["table"] == "Unidades #1"]) == 51  # incl. rows with empty trailing cells

    # evidence: file_read per dashboard read (failed for the unreadable one)
    ev = [e for e in ctx.store.evidence_for(ctx.mission_id) if e.kind == "file_read"]
    assert {Path(e.ref).name for e in ev if e.ok} == {"Dashboard Amara S38.pdf", "Dashboard Amara S39.pdf",
                                                        "Altavista S38.pdf", "Nativa S38.xlsx", "Nativa S39.xlsx"}
    assert [Path(e.ref).name for e in ev if not e.ok] == ["Nativa anexo S39.xlsx"]
    assert all(e.agent_id == "argos" for e in ev)

    # second run: nothing is downloaded again (the broken one is retried), the same fingerprints come back
    before = list(src.downloads)
    ctx2 = run(make_ctx(registry, src, argos_llm()))
    again = run(dash.DashboardsCheck().run(ctx2))
    assert src.downloads[len(before):] == [("m-al-39", "a0")]
    assert {a.fingerprint for a in again.alerts} == {a.fingerprint for a in result.alerts}
    assert "Saved" not in "\n".join(again.notes)


def test_llm_failure_marks_result_not_ok(registry, local):
    llm = FakeLLM()  # no scripted response → LLMError
    result = run(dash.DashboardsCheck().run(run(make_ctx(registry, mailbox(), llm))))
    assert not result.ok  # the engine must not auto-resolve the LLM findings of earlier runs
    assert any("ARGOS comparison failed" in n for n in result.notes)
    assert {a.kind for a in result.alerts} == {"missing_report", "identical_report", "kpi_mismatch"}


def test_without_scope_deterministic_only(registry, local):
    result = run(dash.DashboardsCheck().run(run(make_ctx(registry, mailbox(), None))))
    assert {a.kind for a in result.alerts} == {"missing_report", "identical_report", "kpi_mismatch"}
    assert not result.ok and any("comparison skipped" in n for n in result.notes)


def test_monday_before_new_reports_checks_last_week(registry, local):
    src = mailbox()
    for mid in [m for m in src.messages if m.endswith("-39")]:
        del src.messages[mid]
    monday = datetime(2026, 9, 28, 15, 0, tzinfo=UTC)  # week 40, nothing received yet
    result = run(dash.DashboardsCheck().run(run(make_ctx(registry, src, None, now=monday))))
    assert not [a for a in result.alerts if a.kind == "missing_report"]
    assert any("No dashboards for 2026-W40 yet; checking 2026-W38" in n for n in result.notes)


def test_not_configured(registry, local):
    ctx = run(make_ctx(registry, None, None))
    with pytest.raises(CheckNotConfigured) as exc:
        run(dash.DashboardsCheck().run(ctx))
    assert exc.value.hint == "Connect Outlook in Follow-ups first"

    ctx.mail = AttSource(connected=False)
    with pytest.raises(CheckNotConfigured) as exc:
        run(dash.DashboardsCheck().run(ctx))
    assert exc.value.hint == "Connect Outlook in Follow-ups first"

    class NoAttachments:
        name = "old"

    ctx.mail = NoAttachments()
    with pytest.raises(CheckNotConfigured):
        run(dash.DashboardsCheck().run(ctx))


def test_registered_with_engine():
    engine = pytest.importorskip("atlas.argos.engine")
    assert "dashboards" in engine.registered_checks()
    assert engine.registered_checks()["dashboards"]().name == "dashboards"


# ---------------------------------------------------------------------------
# Units
# ---------------------------------------------------------------------------


def test_week_for_tokens_and_dates():
    assert dash.week_for(W39_AT) == W39
    assert dash.week_for(W39_AT, "Reporte S38.pdf") == W38
    assert dash.week_for(W39_AT, "reporte", "Dashboard Semana 37") == "2026-W37"
    assert dash.week_for(W39_AT, "Sem. 36") == "2026-W36"
    assert dash.week_for(datetime(2027, 1, 4, 18, tzinfo=UTC), "S52") == "2026-W52"  # previous year
    assert dash.week_for(W39_AT, "Torre S2B", "los 38 departamentos") == W39  # not week tokens
    assert dash.shift_week("2026-W01", -1) == "2025-W52"


def test_parse_dates():
    assert dash.parse_dates("30-sep", 2026) == ["2026-09-30"]
    assert dash.parse_dates("30 de octubre de 2026", 2026) == ["2026-10-30"]
    assert dash.parse_dates("2026-12-15", 2026) == ["2026-12-15"]
    assert dash.parse_dates("05/11/2026", 2026) == ["2026-11-05"]
    assert dash.parse_dates("22/51", 2026) == [] and dash.parse_dates("64%", 2026) == []


def test_normalize_pdf_style_text():
    text = ("--- Page 1 ---\nDashboard Nativa\nEscrituradas: 3/4\nFecha de corte: 18-sep-2026\n[table]\n"
            "Unidad | Estatus\nN-1 | Escriturado\nN-2 | Escriturado\nN-3 | Vendido\nN-4 | Disponible")
    snap = dash.normalize(text, year=2026)
    assert snap["kpis"] == {"Escrituradas": "3/4", "Fecha de corte": "18-sep-2026"}
    assert {"label": "Fecha de corte", "date": "2026-09-18", "raw": "18-sep-2026"} in snap["dates"]
    assert [(r["key"], r["cells"]["Estatus"]) for r in snap["rows"]] == [
        ("N-1", "Escriturado"), ("N-2", "Escriturado"), ("N-3", "Vendido"), ("N-4", "Disponible")]
    [mm] = dash.kpi_mismatches(dash.normalize(text.replace("3/4", "1/4"), year=2026))
    assert (mm["kpi"], mm["count"]) == (1, 2)
    assert snap["text_hash"] == dash.text_hash(text.replace("Dashboard", "Dashboard  "))


def test_kpi_mismatch_only_when_confident():
    rows = "\n".join(f"U{i} | {'Escriturada' if i < 5 else 'Vendida'} | x" for i in range(8))
    text = f"Unidades escrituradas: 4/8\n[table]\nUnidad | Estatus | Nota\n{rows}"
    [mm] = dash.kpi_mismatches(dash.normalize(text, year=2026))
    assert (mm["kpi"], mm["count"], mm["column"]) == (4, 5, "Estatus")
    assert dash.kpi_mismatches(dash.normalize(text.replace("4/8", "5/8"), year=2026)) == []
    # the status appears in two columns → not confident
    ambiguous = "Escrituradas: 1\n[table]\nA | B | C\n" + "\n".join(f"u{i} | Escriturada | Escriturada"
                                                                        for i in range(3))
    assert dash.kpi_mismatches(dash.normalize(ambiguous, year=2026)) == []


def test_quote_ok():
    text = "Hito | Fecha\nCimentación | 30-sep-2026 | En proceso"
    assert dash.quote_ok("Cimentación | 30-sep-2026", text)
    assert dash.quote_ok("Cimentación 30-sep-2026  En proceso", text)
    assert not dash.quote_ok("Cimentación | 30-oct-2026", text)
    assert not dash.quote_ok("", text)


def test_state_helpers_keep_other_keys(local):
    path = dash.state_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({"runs": {"last_watch": "x"}}), encoding="utf-8")
    st = dash.load_state()
    assert st["attachments"] == {}
    st["attachments"]["m:a"] = {"sha256": "0" * 64}
    dash.save_state(st)
    data = json.loads(path.read_text(encoding="utf-8"))
    assert data["runs"] == {"last_watch": "x"} and data["dashboards"]["attachments"]["m:a"]["sha256"] == "0" * 64
    path.write_text("{broken", encoding="utf-8")
    assert dash.load_state() == {"attachments": {}}


def test_config_parsing():
    cfg = dash.DashboardsConfig.from_watch({}, digest_projects={"Nativa": ["nativa"]})
    assert cfg.lookback_days == 10 and cfg.extensions == [".pdf", ".xlsx", ".xlsm"]
    assert cfg.matches("Reporte SEMANAL") and cfg.matches("Avance S38.pdf") and not cfg.matches("Factura")
    assert cfg.project_for("NATIVA s38.xlsx") == "Nativa"
    cfg = dash.DashboardsConfig.from_watch({"dashboards": {
        "lookback_days": "x", "match": ["informe", "(bad", ".csv"], "projects": {"Amāra": ["amara"]},
        "expected": "Amāra"}}, digest_projects={"Nativa": ["nativa"]})
    assert cfg.lookback_days == 10 and cfg.extensions == [".csv"] and cfg.projects == {"Amāra": ["amara"]}
    assert cfg.expected == ["Amāra"] and cfg.matches("(bad") and cfg.matches("Informe S38")
    assert len(cfg.notes) == 2
    assert cfg.project_for("Dashboard AMĀRA") == "Amāra"


# ---------------------------------------------------------------------------
# Attachment support in the sources
# ---------------------------------------------------------------------------


class _Msal:
    def get_accounts(self):
        return [{"username": "jd@empresa.com"}]

    def acquire_token_silent(self, scopes, account=None):
        return {"access_token": "tok", "scope": "Mail.Read Mail.ReadWrite"}


def test_graph_attachments():
    def handler(req: httpx.Request) -> httpx.Response:
        path = req.url.path
        if path.endswith("/messages/m1/attachments") and "skip" not in str(req.url):
            assert "isInline" in req.url.params["$select"]
            return httpx.Response(200, json={"value": [
                {"@odata.type": "#microsoft.graph.fileAttachment", "id": "f1", "name": "Nativa S39.xlsx",
                 "size": 1234, "contentType": "application/vnd.ms-excel", "isInline": False},
                {"@odata.type": "#microsoft.graph.fileAttachment", "id": "img", "name": "logo.png", "isInline": True},
                {"@odata.type": "#microsoft.graph.itemAttachment", "id": "it", "name": "RE: algo"},
                {"@odata.type": "#microsoft.graph.referenceAttachment", "id": "ref", "name": "link.pdf"},
            ], "@odata.nextLink": "https://graph.microsoft.com/v1.0/me/messages/m1/attachments?$skip=4"})
        if path.endswith("/messages/m1/attachments"):
            return httpx.Response(200, json={"value": [
                {"@odata.type": "#microsoft.graph.fileAttachment", "id": "f2", "name": "b.pdf", "size": 3}]})
        if path.endswith("/messages/m1/attachments/f1/$value"):
            return httpx.Response(200, content=b"PK\x03\x04data")
        if path.endswith("/mailFolders/inbox/messages"):
            return httpx.Response(200, json={"value": [{
                "id": "m1", "subject": "S39", "receivedDateTime": "2026-09-23T10:00:00Z", "hasAttachments": True,
                "from": {"emailAddress": {"address": "a@b.c"}}}]})
        return httpx.Response(200, json={"value": []})

    src = GraphSource("client", "organizations", app=_Msal(), transport=httpx.MockTransport(handler))
    assert isinstance(src, AttachmentSource)
    metas = run(src.list_attachments("m1"))
    assert [(m.id, m.name, m.size) for m in metas] == [("f1", "Nativa S39.xlsx", 1234), ("f2", "b.pdf", 3)]
    assert run(src.download_attachment("m1", "f1")) == b"PK\x03\x04data"
    [msg] = run(src.list_messages(datetime(2026, 9, 1, tzinfo=UTC)))
    assert msg.has_attachments


def test_folder_eml_and_json_attachments(tmp_path):
    em = EmailMessage()
    em["From"], em["To"], em["Subject"] = "Gerente <g@c.mx>", "jd@empresa.com", "Reporte semanal S39"
    em["Date"] = "Wed, 23 Sep 2026 10:00:00 -0600"
    em.set_content("Adjunto el reporte.")
    em.add_attachment(b"%PDF-1.4 data", maintype="application", subtype="pdf", filename="Nativa S39.pdf")
    em.add_attachment(b"\x89PNG", maintype="image", subtype="png", cid="<logo>", disposition="inline",
                      filename="logo.png")  # an inline image
    (tmp_path / "reporte.eml").write_bytes(bytes(em))
    plain = EmailMessage()
    plain["From"], plain["Subject"], plain["Date"] = "x@y.z", "Hola", "Wed, 23 Sep 2026 10:00:00 -0600"
    plain.set_content("sin adjuntos")
    (tmp_path / "hola.eml").write_bytes(bytes(plain))

    (tmp_path / "pa.json").write_text(json.dumps({
        "Id": "pa1", "Subject": "Dashboard Amara", "From": "g@c.mx", "DateTimeReceived": "2026-09-23T10:00:00Z",
        "Body": "ver adjunto", "HasAttachments": True,
        "Attachments": [{"Name": "Amara S39.pdf", "ContentBytes": base64.b64encode(b"%PDF amara").decode()},
                        {"Name": "logo.png", "ContentBytes": "AAAA", "IsInline": True}],
    }), encoding="utf-8")
    (tmp_path / "pa2.json").write_text(json.dumps({
        "Id": "pa2", "Subject": "Reporte Altavista", "From": "g@c.mx", "DateTimeReceived": "2026-09-23T10:00:00Z",
        "Body": "ver adjunto"}), encoding="utf-8")
    (tmp_path / "pa2_attachments").mkdir()
    (tmp_path / "pa2_attachments" / "Altavista S39.xlsx").write_bytes(b"PK altavista")

    src = FolderSource(tmp_path, me=set())
    assert isinstance(src, AttachmentSource)
    msgs = {m.subject: m for m in run(src.list_messages(datetime(2026, 9, 1, tzinfo=UTC)))}
    assert msgs["Reporte semanal S39"].has_attachments and not msgs["Hola"].has_attachments
    assert msgs["Dashboard Amara"].has_attachments and msgs["Reporte Altavista"].has_attachments

    eml_id = msgs["Reporte semanal S39"].id
    [meta] = run(src.list_attachments(eml_id))  # the inline image is skipped
    assert meta.name == "Nativa S39.pdf" and meta.content_type == "application/pdf"
    assert run(src.download_attachment(eml_id, meta.id)) == b"%PDF-1.4 data"
    assert run(src.list_attachments(msgs["Hola"].id)) == []

    [pa] = run(src.list_attachments("pa1"))
    assert pa.name == "Amara S39.pdf" and run(src.download_attachment("pa1", pa.id)) == b"%PDF amara"
    [sib] = run(src.list_attachments("pa2"))
    assert sib.name == "Altavista S39.xlsx" and run(src.download_attachment("pa2", sib.id)) == b"PK altavista"
    with pytest.raises(KeyError):
        run(src.download_attachment("pa2", "nope"))
    with pytest.raises(KeyError):
        run(src.list_attachments("missing-id"))
