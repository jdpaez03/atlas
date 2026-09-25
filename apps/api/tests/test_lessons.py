"""Lessons (docs/LESSONS.md): comments → proposed rules → approved lessons injected into each agent's prompt;
the private prompt override; SCRIBE's style knobs."""

from __future__ import annotations

import asyncio
import copy
import json
from pathlib import Path

from fastapi.testclient import TestClient
from pptx import Presentation
from pypdf import PdfReader

from atlas.core import paths
from atlas.core.events import EventBus
from atlas.core.registry import AgentRegistry
from atlas.core.store import WorldStore, fold
from atlas.live import AgentLoader, FakeLLM, LiveConfig, LiveEngine, ModelConfig, NodeContext, PriceTable
from atlas.live.lessons import (
    LESSONS_HEADER,
    LessonBook,
    lessons_block,
    lessons_path,
    load_file,
    with_lessons,
)
from atlas.live.llm import call_text, tool_use
from atlas.publish.brand import load_brand
from atlas.publish.deck import render_deck
from atlas.publish.pdf import DocMeta, render_pdf
from atlas.publish.spec import DEFAULT_STYLE, normalize, normalize_style

MODELS = ModelConfig(orchestrator="o", default="s", fast="f")


def _store() -> WorldStore:
    return WorldStore(AgentRegistry.load(), EventBus())


def test_direct_lesson_is_active_persisted_and_injected():
    store = _store()
    book = LessonBook(store)
    initial = store.snapshot()

    async def go():
        a = await book.direct("scribe", "  Limita las tablas del deck a 8 filas.  ")
        b = await book.direct("*", "Escribe en español de México.")
        c = await book.direct("oracle", "Incluye siempre un escenario bajista.", node="personal")
        return a, b, c

    a, b, c = asyncio.run(go())
    assert a.status == "active" and a.origin == "direct" and a.text == "Limita las tablas del deck a 8 filas."
    on_disk = json.loads(lessons_path().read_text(encoding="utf-8"))["lessons"]
    assert [x["id"] for x in on_disk] == [a.id, b.id, c.id]
    block = lessons_block(store, "scribe", "corporate")
    assert block.startswith(LESSONS_HEADER) and "- Limita las tablas" in block and "- Escribe en español" in block
    assert "escenario bajista" not in lessons_block(store, "oracle", "corporate")  # other node
    assert "escenario bajista" in lessons_block(store, "oracle", "personal")
    assert with_lessons(["role"], _store(), "scribe") == ["role"]  # nothing to add
    assert fold(initial, store.bus.history()).lessons == store.snapshot().lessons
    # a new store/book (server restart) reloads them from the file
    again = _store()
    LessonBook(again)
    assert [x.id for x in again.lessons()] == [a.id, b.id, c.id]


def _distill_llm(rules: list[str], replaces: list[str] | None = None) -> FakeLLM:
    llm = FakeLLM()
    llm.when(lambda kw: "propose_lessons" in {t["name"] for t in kw.get("tools", [])},
             tool_use("propose_lessons", {"rules": rules, "replaces": replaces or []}), repeat=True)
    return llm


def test_comment_becomes_proposals_and_approval_retires_the_replaced_lesson():
    store = _store()
    book = LessonBook(store)

    async def go():
        old = await book.direct("scribe", "Usa 12 filas por tabla.")
        llm = _distill_llm(["Limita las tablas del deck a 6 filas por lámina."], replaces=[old.id, "lsn_bogus"])
        engine = LiveEngine(store, llm=llm, config=LiveConfig(models=MODELS), prices=PriceTable())
        props = await book.comment(engine, "scribe", "las tablas del deck son eternas")
        call = llm.calls[0]
        assert "SCRIBE" in call_text(call) and "las tablas del deck son eternas" in call_text(call)
        assert f"[{old.id}] Usa 12 filas" in call_text(call)  # it sees the current lessons
        assert call["model"] == "f"  # the fast model
        approved = await book.decide(props[0].id, status="active")
        return old, props, approved

    old, props, approved = asyncio.run(go())
    assert len(props) == 1 and props[0].status == "proposed" and props[0].replaces == [old.id]
    assert props[0].comment == "las tablas del deck son eternas"
    assert approved.status == "active" and approved.decided_at is not None
    assert store.lesson(old.id).status == "retired"
    block = lessons_block(store, "scribe")
    assert "6 filas" in block and "12 filas" not in block
    assert store.mission_reports_for  # no mission was created for the distillation
    assert store.snapshot().missions == []


def test_comment_without_a_backend_keeps_the_comment_as_the_proposal(monkeypatch):
    monkeypatch.setenv("ATLAS_LLM_BACKEND", "none")
    store = _store()
    book = LessonBook(store)

    class NoBackend:
        def backend_info(self):
            from atlas.live.backend import BackendInfo

            return BackendInfo(None)

    props = asyncio.run(book.comment(NoBackend(), "oracle", "  pon siempre el escenario bajista "))
    assert [p.text for p in props] == ["pon siempre el escenario bajista"] and props[0].status == "proposed"
    assert lessons_block(store, "oracle") == ""  # not active until approved


def test_lessons_reach_the_agents_in_a_live_mission():
    reg = AgentRegistry.load()
    store = WorldStore(reg, EventBus())
    book = LessonBook(store)
    plan = [{"ref": "a", "title": "Analyze", "description": "x", "assigned_to": "oracle"},
            {"ref": "b", "title": "Package", "description": "y", "assigned_to": "alfred"}]
    rep = {"findings": [{"kind": "FACT", "statement": "ok", "confidence": "HIGH"}], "confidence": "HIGH"}
    llm = FakeLLM()
    llm.when(lambda kw: "STAGE: PLANNING" in call_text(kw), tool_use("create_plan", {"tasks": plan}))
    llm.when(lambda kw: "Your task: Analyze" in call_text(kw), tool_use("submit_report", {**rep, "asked_to": "a"}))
    llm.when(lambda kw: "Your task: Package" in call_text(kw), tool_use("submit_report", {**rep, "asked_to": "b"}))
    llm.when(lambda kw: "STAGE: REVIEW" in call_text(kw), tool_use("request_followups", {"tasks": []}))
    llm.when(lambda kw: "STAGE: CONSOLIDATION" in call_text(kw), tool_use("submit_mission_report", {
        "executive_summary": "ok", "objective_status": "ACHIEVED",
        "key_findings": [{"kind": "FACT", "statement": "ok", "confidence": "HIGH"}]}))

    async def go():
        await book.direct("oracle", "Incluye siempre un escenario bajista.")
        await book.direct("*", "Responde en español de México.")
        await book.direct("atlas", "Planea máximo 4 tareas.")
        prop = (await book.comment(None, "alfred", "firma como Dirección"))[0]  # proposed only: not injected
        engine = LiveEngine(store, loader=AgentLoader(reg, MODELS),
                            context=NodeContext(reg, local_dir=Path("/nonexistent")),
                            config=LiveConfig(models=MODELS, max_turns=3, audit=False, publish=False), llm=llm,
                            prices=PriceTable())
        m = await engine.start("Evaluate", "corporate")
        await asyncio.wait_for(engine.wait(m.id), 10)
        return prop

    prop = asyncio.run(go())
    assert prop.text == "firma como Dirección"
    oracle = next(c for c in llm.calls if "Your task: Analyze" in call_text(c))
    alfred = next(c for c in llm.calls if "Your task: Package" in call_text(c))
    plan_call = next(c for c in llm.calls if "STAGE: PLANNING" in call_text(c))
    assert "escenario bajista" in call_text(oracle) and "español de México" in call_text(oracle)
    assert "escenario bajista" not in call_text(alfred) and "español de México" in call_text(alfred)
    assert "firma como Dirección" not in call_text(alfred)
    assert "Planea máximo 4 tareas" in call_text(plan_call)


def test_http_lessons(monkeypatch):
    from atlas.main import app

    with TestClient(app) as client:
        assert client.get("/lessons").json() == []
        created = client.post("/lessons", json={"agent_id": "scribe", "text": "Agrega índice al PDF", "mode": "direct"})
        assert created.status_code == 200 and created.json()[0]["status"] == "active"
        lid = created.json()[0]["id"]
        proposed = client.post("/lessons", json={"agent_id": "oracle", "text": "más escenarios"}).json()
        assert proposed[0]["status"] == "proposed"  # no backend in tests → the comment itself, for approval
        assert client.patch(f"/lessons/{proposed[0]['id']}", json={"status": "dismissed"}).json()["status"] == "dismissed"
        edited = client.patch(f"/lessons/{lid}", json={"text": "Agrega índice al PDF de estudios"}).json()
        assert edited["text"] == "Agrega índice al PDF de estudios" and edited["status"] == "active"
        assert [x["id"] for x in client.get("/lessons", params={"status": "active"}).json()] == [lid]
        assert client.post("/lessons", json={"agent_id": "nobody", "text": "x"}).status_code == 404
        assert client.patch("/lessons/lsn_nope", json={"status": "active"}).status_code == 404
        assert client.patch(f"/lessons/{lid}", json={"status": "bogus"}).status_code == 422
        assert any(e["type"] == "lesson.upserted" for e in client.get("/events").json())
    assert next(x.id for x in load_file()) == lid


def test_private_prompt_overrides_the_public_one():
    reg = AgentRegistry.load()
    loader = AgentLoader(reg, MODELS)
    assert loader.resolve("scribe").source == "yaml"
    p = paths.local_dir() / "agents" / "scribe.md"
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text("---\nname: scribe\n---\n# SCRIBE privado\nEscribe como PAGA.\n", encoding="utf-8")
    r = loader.resolve("scribe")
    assert r.source == "local" and r.role_prompt.startswith("# SCRIBE privado") and "name:" not in r.role_prompt
    p.write_text("", encoding="utf-8")
    assert loader.resolve("scribe").source == "yaml"


# ---------------------------------------------------------------------------
# SCRIBE style knobs
# ---------------------------------------------------------------------------

SPEC = {
    "title": "Oferta *competidora*", "summary": ["Resumen."],
    "sections": [{"title": "Uno", "blocks": [{"type": "paragraph", "text": "a"}]},
                 {"title": "Dos", "blocks": [{"type": "chart", "chart": {
                     "chart_type": "bar", "categories": ["A", "B"], "series": [{"name": "s", "values": [1, 2]}]}}]}],
    "slides": [{"type": "section", "title": "Uno"},
               {"type": "table", "title": "Tabla", "table": {"columns": ["P", "V"],
                                                             "rows": [[f"P{i}", i] for i in range(11)]}},
               {"type": "section", "title": "Dos"},
               {"type": "chart", "title": "G", "chart": {"chart_type": "bar", "categories": ["A", "B"],
                                                        "series": [{"name": "s", "values": [1, 2]}]}}],
}


def test_style_normalization():
    assert normalize_style(None) == DEFAULT_STYLE
    assert normalize_style({"deck_density": "huge", "table_rows_per_slide": 99, "toc": "yes", "agenda": True}) == {
        **DEFAULT_STYLE, "table_rows_per_slide": 14, "agenda": True}
    spec, errors = normalize(copy.deepcopy(SPEC))
    assert errors == [] and spec["style"] == DEFAULT_STYLE
    assert len(spec["slides"][1]["table"]["rows"]) == 11  # deck tables are no longer cut at 12 → split instead


def test_style_knobs_change_the_documents(tmp_path):
    brand = load_brand("corporate")
    base, _ = normalize(copy.deepcopy(SPEC))
    styled, _ = normalize({**copy.deepcopy(SPEC), "style": {
        "agenda": True, "toc": True, "table_rows_per_slide": 4, "chart_data_labels": True, "section_numbers": False,
        "deck_density": "compact"}})
    plain_deck = Presentation(str(render_deck(base, brand, "d", [], tmp_path / "a.pptx")))
    deck = Presentation(str(render_deck(styled, brand, "d", [], tmp_path / "b.pptx")))
    # default: 11 rows at 8/slide = 2 table slides; styled: agenda + 11 rows at 4/slide = 3 table slides
    assert len(plain_deck.slides) == 2 + 4 + 1
    assert len(deck.slides) == 2 + 4 + 1 + 1 + 1
    texts = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in deck.slides]
    assert "Agenda" in texts[1] and "01" in texts[1] and "Dos" in texts[1]
    assert sum("Tabla (cont.)" in t for t in texts) == 2
    chart = next(sh.chart for s in deck.slides for sh in s.shapes if sh.has_chart)
    assert chart.plots[0].has_data_labels

    a = render_pdf(base, brand, DocMeta("d"), tmp_path / "a.pdf")
    b = render_pdf(styled, brand, DocMeta("d"), tmp_path / "b.pdf")
    ta = " ".join(p.extract_text() for p in PdfReader(str(a)).pages)
    tb = " ".join(p.extract_text() for p in PdfReader(str(b)).pages)
    assert "Contenido" not in ta and "Contenido" in tb and "01" in ta and "01" not in tb.replace("2026", "")
    assert len(PdfReader(str(b)).pages) == len(PdfReader(str(a)).pages) + 1  # the contents page
