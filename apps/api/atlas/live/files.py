"""File access for live agents (docs/PHASE3.md § A): sandbox policy, tools, text extraction, deliverables.

Agents read the user's files ONLY through these tools, and only inside the read roots of their mission's node:

    ATLAS_FILE_ROOTS_<NODE>   read roots, separated by ';' (default: corporate -> ~/Documents, others -> none)
    + the mission's attachments folder and its own outputs folder (always readable)

`FileSandbox.resolve()` is the single policy function: real path (symlinks followed) inside an allowed root, the
deepest matching root must belong to this node (a node never reads another node's roots), nothing on the deny
list (.git, .env*, *.pem, *.key, id_rsa*, .ssh, ~/.claude, ...) and, inside ATLAS_LOCAL_DIR, only this node's
context, this mission's attachments and this mission's outputs.

Reading is read-only and bounded: 25 MB per file, text paged by `offset` in chunks of at most
ATLAS_FILE_MAX_CHARS (default 60k). Listing and searching are capped (entries, depth, visited directories) and
skip heavy folders (node_modules, .git, virtualenvs...).

`write_deliverable` writes ONLY to outputs/<node>/<mission_id>/, never overwriting ("name (2).ext").

Everything here is synchronous and store-free: `FileTools.run(name, args)` returns a `FileResult` (tool text +
what to record as evidence). The runtime calls it in a worker thread and records the evidence.
"""

from __future__ import annotations

import csv
import fnmatch
import inspect
import io
import json
import os
import re
import threading
from collections import OrderedDict
from collections.abc import Callable, Iterator
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from urllib.parse import quote

from ..core import paths
from ..core.models import Attachment

MAX_FILE_BYTES = 25 * 1024 * 1024
DEFAULT_MAX_CHARS = 60_000
EXTRACT_LIMIT = 2_000_000  # stop extracting a huge file after this many chars
MAX_LIST = 500
MAX_HITS = 100
MAX_DEPTH = 6  # recursive listing / search depth below the start folder
MAX_VISITED = 5_000  # directories scanned per list/search call
MAX_WRITE_CHARS = 5_000_000
FORMATS = ("md", "txt", "csv", "json", "xlsx", "docx")
SKIP_DIRS = {"node_modules", ".git", "__pycache__", ".venv", "venv", ".tox", ".cache", ".next", ".mypy_cache",
             ".pytest_cache", ".ruff_cache", "site-packages", "$recycle.bin", "system volume information"}
DENY_DIRS = {".git", ".ssh", ".gnupg", ".aws", ".claude"}
TEXT_EXT = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl", ".yaml", ".yml", ".xml", ".html", ".htm",
    ".log", ".ini", ".cfg", ".toml", ".rst", ".sql", ".py", ".js", ".ts", ".tsx", ".jsx", ".css", ".sh",
    ".ps1", ".bat", ".java", ".c", ".h", ".cpp", ".cs", ".go", ".rs", ".rb", ".php", ".r", ".tex", ".srt",
}
_DRIVE = re.compile(r"^[A-Za-z]:[\\/]")


class FileAccessError(Exception):
    """A path the sandbox refuses (the message is safe to show the agent)."""


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------


def env_key(node: str) -> str:
    return "ATLAS_FILE_ROOTS_" + re.sub(r"[^A-Z0-9]", "_", node.upper())


def max_chars() -> int:
    try:
        return max(1000, int(os.getenv("ATLAS_FILE_MAX_CHARS", DEFAULT_MAX_CHARS)))
    except ValueError:
        return DEFAULT_MAX_CHARS


def _split_roots(raw: str) -> list[Path]:
    out = []
    for part in raw.split(";"):
        part = part.strip().strip('"')
        if part:
            out.append(Path(os.path.expanduser(_seps(part))))
    return out


def node_roots(node: str) -> list[Path]:
    """Configured read roots of a node (unresolved): ATLAS_FILE_ROOTS_<NODE>, default corporate -> ~/Documents."""
    raw = os.environ.get(env_key(node))
    if raw is None:
        return [Path.home() / "Documents"] if node == "corporate" else []
    return _split_roots(raw)


def _other_node_roots(node: str) -> list[tuple[Path, str]]:
    """Roots configured for every other node (from the environment, plus the corporate default)."""
    prefix = "ATLAS_FILE_ROOTS_"
    mine = env_key(node)
    out: list[tuple[Path, str]] = []
    keys = {k for k in os.environ if k.startswith(prefix)}
    if env_key("corporate") not in keys:
        keys.add(env_key("corporate"))
    for key in sorted(keys):
        if key == mine:
            continue
        owner = key[len(prefix):].lower()
        roots = node_roots("corporate") if key == env_key("corporate") else _split_roots(os.environ.get(key, ""))
        out += [(r, owner) for r in roots]
    return out


def _seps(raw: str) -> str:
    """Accept Windows separators everywhere (on POSIX a backslash becomes '/')."""
    return raw.replace("\\", "/") if os.sep == "/" else raw


def _real(p: Path) -> Path:
    try:
        return p.resolve()
    except (OSError, RuntimeError):  # symlink loop, permission...
        return Path(os.path.abspath(p))


def _within(path: Path, root: Path) -> bool:
    return path == root or path.is_relative_to(root)


def _denied_part(part: str) -> bool:
    p = part.lower()
    return (p in DENY_DIRS or p.endswith((".pem", ".key"))
            or p.startswith((".env", "id_rsa", "id_ed25519", "id_ecdsa", "id_dsa")))


# ---------------------------------------------------------------------------
# Sandbox policy
# ---------------------------------------------------------------------------


@dataclass
class FileSandbox:
    node: str
    mission_id: str
    roots: list[Path]  # resolved read roots of this node (configured roots + attachments + outputs)
    others: list[tuple[Path, str]]  # resolved roots of other nodes
    local: Path
    outputs: Path
    attachments: Path
    home_claude: Path
    max_chars: int = DEFAULT_MAX_CHARS

    @classmethod
    def for_mission(cls, node: str, mission_id: str) -> FileSandbox:
        local = _real(paths.local_dir())
        attachments = _real(paths.attachments_dir(mission_id))
        outputs = _real(paths.outputs_dir(node, mission_id))
        roots = [_real(r) for r in node_roots(node)]
        return cls(
            node=node, mission_id=mission_id,
            roots=[*roots, attachments, outputs],
            others=[(_real(r), owner) for r, owner in _other_node_roots(node)],
            local=local, outputs=outputs, attachments=attachments,
            home_claude=_real(Path.home() / ".claude"), max_chars=max_chars(),
        )

    @property
    def user_roots(self) -> list[Path]:
        return [r for r in self.roots if r not in (self.attachments, self.outputs)]

    def check(self, path: Path, *, traverse: bool = False) -> str | None:
        """Why the (already resolved) path is not readable, or None when it is. With `traverse`, the ATLAS
        folders that only lead to an allowed area (atlas-local, atlas-local/context...) pass too, so a walk
        can reach this node's context; their other entries are still checked one by one."""
        if any(_denied_part(p) for p in path.parts) or _within(path, self.home_claude):
            return "access to this path is denied (protected file or folder)"
        if _within(path, self.local):
            fold = (lambda x: x.lower()) if os.name == "nt" else (lambda x: x)
            rel = [fold(p) for p in path.relative_to(self.local).parts]
            node, mid = fold(self.node), fold(self.mission_id)
            allowed = (["context", node], ["missions", mid, "attachments"], ["outputs", node, mid])
            ok = any(rel[:len(a)] == a for a in allowed)
            if not ok and traverse:
                ok = any(len(rel) < len(a) and rel == a[:len(rel)] for a in allowed)
            if not ok:
                return "access to this ATLAS folder is denied"
        matches = [(len(r.parts), True) for r in self.roots if _within(path, r)]
        matches += [(len(r.parts), False) for r, _ in self.others if _within(path, r)]
        if not any(mine for _, mine in matches):
            return "the path is outside the allowed folders"
        if not max(matches)[1]:  # deepest root (ties go to this node) belongs to another node
            return "the path belongs to another node's folders"
        return None

    def resolve(self, raw: str | None, *, must_exist: bool = True) -> Path:
        """Agent-supplied path -> real, allowed path. Relative paths are tried against each root in order."""
        text = (raw or "").strip().strip('"').strip("'")
        if not text:
            raise FileAccessError("no path given")
        if "\x00" in text:
            raise FileAccessError("invalid path")
        if os.name != "nt" and _DRIVE.match(text):
            raise FileAccessError(f"'{raw}' is outside the allowed folders")
        text = os.path.expanduser(_seps(text))
        p = Path(text)
        if p.is_absolute():
            candidates = [p]
        else:
            candidates = [r / p for r in self.roots]
            existing = [c for c in candidates if c.exists()]
            candidates = existing or candidates[:1]
        if not candidates:
            raise FileAccessError("no folders are readable in this node")
        real = _real(candidates[0])
        reason = self.check(real)
        if reason:
            raise FileAccessError(f"'{raw}': {reason}")
        if must_exist and not real.exists():
            raise FileAccessError(f"'{raw}' does not exist")
        return real


# ---------------------------------------------------------------------------
# Text extraction (cached per path + mtime + size, so paging a PDF doesn't re-parse it)
# ---------------------------------------------------------------------------

_CACHE: OrderedDict[tuple[Any, ...], tuple[str, str]] = OrderedDict()
_CACHE_LOCK = threading.Lock()
_CACHE_SIZE = 16


class _Sink:
    """Collects extracted text up to EXTRACT_LIMIT chars."""

    def __init__(self) -> None:
        self.parts: list[str] = []
        self.n = 0
        self.full = False

    def add(self, line: str) -> bool:
        if self.full:
            return False
        self.parts.append(line)
        self.n += len(line) + 1
        if self.n >= EXTRACT_LIMIT:
            self.full = True
            self.parts.append(f"[extraction stopped after {EXTRACT_LIMIT:,} characters]")
        return not self.full

    def text(self) -> str:
        return "\n".join(self.parts)


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.isoformat(sep=" ", timespec="minutes") if (v.hour or v.minute) else v.date().isoformat()
    if isinstance(v, (date, time)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return " ".join(str(v).split())


def decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _text(path: Path, sheet: str | None) -> tuple[str, str]:
    data = path.read_bytes()
    if b"\x00" in data[:4096]:
        raise FileAccessError(f"'{path.name}' is a binary file; no text can be extracted from it")
    return decode_text(data), "text"


def _pdf(path: Path, sheet: str | None) -> tuple[str, str]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise FileAccessError(f"'{path.name}' is password-protected") from exc
    sink = _Sink()
    for i, page in enumerate(reader.pages, 1):
        if not sink.add(f"--- Page {i} ---"):
            break
        if not sink.add((page.extract_text() or "").strip()):
            break
    text = sink.text()
    if not re.sub(r"--- Page \d+ ---|\s", "", text):
        text += "\n[no extractable text: the PDF may be scanned images]"
    return text, f"pdf, {len(reader.pages)} pages"


def _xlsx(path: Path, sheet: str | None) -> tuple[str, str]:
    import openpyxl

    values = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(str(path), read_only=True, data_only=False)
    try:
        names = values.sheetnames
        chosen = names
        if sheet:
            match = [n for n in names if n.lower() == sheet.strip().lower()]
            if not match:
                raise FileAccessError(f"no sheet '{sheet}' in {path.name}; sheets: {', '.join(names)}")
            chosen = match
        sink = _Sink()
        for name in chosen:
            if not sink.add(f"## Sheet: {name}"):
                break
            ws_v, ws_f = values[name], formulas[name]
            for row_v, row_f in zip(ws_v.iter_rows(), ws_f.iter_rows(), strict=False):
                cells = []
                for cv, cf in zip(row_v, row_f, strict=False):
                    v = getattr(cv, "value", None)
                    if v is None:
                        f = getattr(cf, "value", None)
                        v = f if isinstance(f, str) and f.startswith("=") else None
                    cells.append(_cell(v))
                while cells and not cells[-1]:
                    cells.pop()
                if cells and not sink.add(" | ".join(cells)):
                    break
            if sink.full:
                break
            sink.add("")
        return sink.text().rstrip(), f"workbook, sheets: {', '.join(names)}"
    finally:
        values.close()
        formulas.close()


def _docx(path: Path, sheet: str | None) -> tuple[str, str]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    sink = _Sink()
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            if style.lower().startswith("heading"):
                level = re.sub(r"\D", "", style) or "1"
                text = "#" * min(6, int(level)) + " " + text
            elif "list" in style.lower():
                text = "- " + text
            if not sink.add(text):
                break
        elif tag == "tbl":
            table = Table(child, document)
            sink.add("")
            for row in table.rows:
                if not sink.add("| " + " | ".join(_cell(c.text) for c in row.cells) + " |"):
                    break
            sink.add("")
    return sink.text().strip(), "word document"


def _pptx(path: Path, sheet: str | None) -> tuple[str, str]:
    from pptx import Presentation

    prs = Presentation(str(path))
    sink = _Sink()
    for i, slide in enumerate(prs.slides, 1):
        if not sink.add(f"--- Slide {i} ---"):
            break
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                sink.add(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    sink.add("| " + " | ".join(_cell(c.text) for c in row.cells) + " |")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                sink.add(f"Notes: {notes}")
    return sink.text().strip(), f"presentation, {len(prs.slides)} slides"


EXTRACTORS: dict[str, Callable[[Path, str | None], tuple[str, str]]] = {
    ".pdf": _pdf, ".xlsx": _xlsx, ".xlsm": _xlsx, ".docx": _docx, ".pptx": _pptx,
}
UNSUPPORTED = {".xls": "old .xls workbooks are not supported (save it as .xlsx)",
               ".doc": "old .doc files are not supported (save it as .docx)",
               ".ppt": "old .ppt files are not supported (save it as .pptx)"}


def extract_text(path: Path, sheet: str | None = None) -> tuple[str, str]:
    """(full text, kind description) of a readable file. Raises FileAccessError."""
    st = path.stat()
    if st.st_size > MAX_FILE_BYTES:
        raise FileAccessError(f"'{path.name}' is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB")
    ext = path.suffix.lower()
    if ext in UNSUPPORTED:
        raise FileAccessError(UNSUPPORTED[ext])
    key = (str(path), st.st_mtime_ns, st.st_size, (sheet or "").lower())
    with _CACHE_LOCK:
        if key in _CACHE:
            _CACHE.move_to_end(key)
            return _CACHE[key]
    extractor = EXTRACTORS.get(ext, _text)
    try:
        out = extractor(path, sheet)
    except FileAccessError:
        raise
    except Exception as exc:
        raise FileAccessError(f"could not read '{path.name}': {type(exc).__name__}: {exc}") from exc
    with _CACHE_LOCK:
        _CACHE[key] = out
        while len(_CACHE) > _CACHE_SIZE:
            _CACHE.popitem(last=False)
    return out


# ---------------------------------------------------------------------------
# Deliverable writers
# ---------------------------------------------------------------------------


def _xlsx_bytes(sheets: dict[str, Any]) -> bytes:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for raw_name, rows in sheets.items():
        name = re.sub(r"[\[\]:*?/\\]", "_", str(raw_name)).strip()[:31] or "Sheet"
        base, n = name, 2
        while name.lower() in used:
            name = f"{base[:28]} {n}"
            n += 1
        used.add(name.lower())
        ws = wb.create_sheet(name)
        for row in rows if isinstance(rows, list) else []:
            cells = row if isinstance(row, list) else [row]
            ws.append([c if isinstance(c, (int, float, bool)) or c is None else str(c) for c in cells])
        for cell in ws[1] if ws.max_row >= 1 and rows else []:
            cell.font = Font(bold=True)
    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _runs(paragraph: Any, text: str) -> None:
    pos = 0
    for m in _BOLD.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        paragraph.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _docx_bytes(content: str) -> bytes:
    """Markdown-ish text -> Word: '#' headings, '-'/'*' bullets, '1.' numbered, '|' tables, paragraphs."""
    import docx

    document = docx.Document()
    para: list[str] = []
    table: list[list[str]] = []

    def flush_para() -> None:
        if para:
            _runs(document.add_paragraph(), " ".join(para))
            para.clear()

    def flush_table() -> None:
        if table:
            cols = max(len(r) for r in table)
            t = document.add_table(rows=len(table), cols=cols)
            t.style = "Table Grid"
            for i, row in enumerate(table):
                for j in range(cols):
                    cell = t.cell(i, j)
                    cell.text = row[j] if j < len(row) else ""
                    if i == 0 and cell.paragraphs[0].runs:
                        cell.paragraphs[0].runs[0].bold = True
            table.clear()

    for raw in content.splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith("|"):
            flush_para()
            if re.fullmatch(r"\|?[\s:|-]+\|?", stripped) and "-" in stripped:
                continue  # markdown separator row
            table.append([c.strip() for c in stripped.strip("|").split("|")])
            continue
        flush_table()
        if not stripped:
            flush_para()
        elif m := re.match(r"^(#{1,6})\s+(.*)$", stripped):
            flush_para()
            document.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4))
        elif m := re.match(r"^[-*+•]\s+(.*)$", stripped):
            flush_para()
            _runs(document.add_paragraph(style="List Bullet"), m.group(1))
        elif m := re.match(r"^\d+[.)]\s+(.*)$", stripped):
            flush_para()
            _runs(document.add_paragraph(style="List Number"), m.group(1))
        else:
            para.append(stripped)
    flush_para()
    flush_table()
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _csv_text(rows: list[Any]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    for row in rows:
        writer.writerow(row if isinstance(row, list) else [row])
    return buf.getvalue()


def render_deliverable(fmt: str, content: str | None, sheets: dict[str, Any] | None) -> bytes:
    """Bytes of a deliverable. Raises FileAccessError on bad input."""
    if fmt not in FORMATS:
        raise FileAccessError(f"unsupported format '{fmt}' (use one of: {', '.join(FORMATS)})")
    if content is not None and not isinstance(content, str):
        content = json.dumps(content, ensure_ascii=False, indent=2) if fmt == "json" else str(content)
    if content and len(content) > MAX_WRITE_CHARS:
        raise FileAccessError(f"content is larger than {MAX_WRITE_CHARS:,} characters")
    if sheets is not None and not isinstance(sheets, dict):
        raise FileAccessError("sheets must be an object {sheet name: rows[][]}")
    if fmt == "xlsx":
        if not sheets:
            if not content:
                raise FileAccessError("xlsx needs sheets ({name: rows[][]}) or CSV content")
            sheets = {"Sheet1": list(csv.reader(io.StringIO(content)))}
        return _xlsx_bytes(sheets)
    if fmt == "csv" and not content and sheets:
        content = _csv_text(next(iter(sheets.values())) or [])
    if not content:
        raise FileAccessError(f"{fmt} needs content")
    if fmt == "json":
        try:
            return (json.dumps(json.loads(content), ensure_ascii=False, indent=2) + "\n").encode("utf-8")
        except json.JSONDecodeError as exc:
            raise FileAccessError(f"content is not valid JSON: {exc}") from exc
    if fmt == "docx":
        return _docx_bytes(content)
    return content.encode("utf-8")


# ---------------------------------------------------------------------------
# Tools
# ---------------------------------------------------------------------------


@dataclass
class FileResult:
    """A tool's text for the model, plus what the runtime records as evidence."""

    text: str
    ok: bool
    kind: str  # Evidence.kind
    ref: str
    detail: str = ""
    attachment: Attachment | None = None
    extra: dict[str, Any] = field(default_factory=dict)


def _size(n: int) -> str:
    if n < 1024:
        return f"{n} B"
    if n < 1024 * 1024:
        return f"{n / 1024:.1f} KB"
    return f"{n / (1024 * 1024):.1f} MB"


def _mtime(ts: float) -> str:
    return datetime.fromtimestamp(ts).astimezone().strftime("%Y-%m-%d %H:%M")  # local time


def _short(text: str, n: int = 60) -> str:
    text = " ".join(str(text).split())
    return text if len(text) <= n else text[: n - 1] + "…"


def _int(value: Any, default: int) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _bool(value: Any) -> bool:
    return value is True or str(value).strip().lower() in ("1", "true", "yes")


def download_url(mission_id: str, name: str) -> str:
    return f"/missions/{quote(mission_id)}/files/outputs/{quote(name)}"


class FileTools:
    NAMES = ("list_files", "search_files", "read_file", "write_deliverable")

    def __init__(self, sandbox: FileSandbox):
        self.sb = sandbox

    # -- activity lines (shown while the tool runs; derived from the tool call only) --

    def activity(self, name: str, args: dict[str, Any]) -> str:
        def base(p: Any) -> str:
            s = _seps(str(p or "")).rstrip("/")
            return s.rsplit("/", 1)[-1] or s

        if name == "read_file":
            return _short(f"Reading {base(args.get('path'))}", 90)
        if name == "list_files":
            return _short(f"Listing {base(args.get('path')) or 'readable folders'}", 90)
        if name == "search_files":
            return _short(f"Searching files · {args.get('query', '')}", 90)
        if name == "write_deliverable":
            fname = paths.safe_name(str(args.get("filename") or "deliverable"))
            fmt = str(args.get("format") or "").lower().lstrip(".")
            return _short(f"Writing {_with_ext(fname, fmt)}", 90)
        return name

    # -- dispatch --

    def run(self, name: str, args: dict[str, Any]) -> FileResult:
        fn = {"list_files": self.list_files, "search_files": self.search_files, "read_file": self.read_file,
              "write_deliverable": self.write_deliverable}[name]
        kind = {"read_file": "file_read", "write_deliverable": "file_written"}.get(name, "file_listed")
        try:
            params = inspect.signature(fn).parameters
            return fn(**{k: v for k, v in args.items() if k in params and v is not None})
        except FileAccessError as exc:
            ref = str(args.get("path") or args.get("filename") or args.get("under") or args.get("query") or "")
            return FileResult(f"Error: {exc}", False, kind, ref, str(exc))
        except OSError as exc:
            ref = str(args.get("path") or args.get("filename") or "")
            return FileResult(f"Error: {exc.strerror or exc}", False, kind, ref, str(exc.strerror or exc))

    # -- list / search --

    def _walk(self, start: Path, recursive: bool) -> Iterator[tuple[Path, os.stat_result, bool, int]]:
        """(real path, stat, is_dir, depth) under `start`, allowed entries only, capped and pruned."""
        stack: list[tuple[Path, int]] = [(start, 0)]
        visited = 0
        while stack and visited < MAX_VISITED:
            folder, depth = stack.pop(0)
            visited += 1
            in_local = _within(folder, self.sb.local)
            try:
                entries = sorted(os.scandir(folder), key=lambda e: e.name.lower())
            except OSError:
                continue
            for entry in entries:
                try:
                    is_link = entry.is_symlink()
                    path = _real(Path(entry.path)) if is_link else folder / entry.name
                    is_dir = entry.is_dir()
                    # files: only links and protected names need the full check; folders always (another
                    # node's root or an ATLAS-private folder may sit inside this root)
                    if (is_dir or is_link or in_local or _denied_part(entry.name)) and self.sb.check(
                            path, traverse=is_dir):
                        continue
                    st = entry.stat()
                except OSError:
                    continue
                if is_dir and (entry.name.lower() in SKIP_DIRS or entry.name.startswith(".")):
                    continue
                yield path, st, is_dir, depth
                if is_dir and recursive and depth + 1 < MAX_DEPTH and not is_link:
                    stack.append((path, depth + 1))

    def _roots_listing(self) -> FileResult:
        lines = ["Readable folders (read-only):"]
        for r in self.sb.user_roots:
            lines.append(f"- {r}" + ("" if r.is_dir() else "  (not found)"))
        if not self.sb.user_roots:
            lines.append("- (none configured for this node)")
        att = self.sb.attachments
        names = sorted(p.name for p in att.iterdir()) if att.is_dir() else []
        lines.append(f"Mission attachments: {att}" + (f" ({len(names)} files)" if names else " (none)"))
        lines += [f"  - {n}" for n in names[:50]]
        lines.append(f"Your deliverables folder: {self.sb.outputs}")
        return FileResult("\n".join(lines), True, "file_listed", "(readable folders)", f"{len(self.sb.roots)} roots")

    def list_files(self, path: str | None = None, pattern: str | None = None, recursive: Any = False) -> FileResult:
        if not (path or "").strip():
            return self._roots_listing()
        folder = self.sb.resolve(path)
        if not folder.is_dir():
            raise FileAccessError(f"'{path}' is not a folder")
        rec = _bool(recursive)
        pat = (pattern or "").strip().lower()
        lines: list[str] = []
        truncated = False
        for p, st, is_dir, _ in self._walk(folder, rec):
            if pat and (is_dir or not fnmatch.fnmatch(p.name.lower(), pat)):
                continue
            rel = p.relative_to(folder).as_posix() if _within(p, folder) else p.name
            lines.append(f"{rel}/\t\t{_mtime(st.st_mtime)}" if is_dir
                         else f"{rel}\t{_size(st.st_size)}\t{_mtime(st.st_mtime)}")
            if len(lines) >= MAX_LIST:
                truncated = True
                break
        head = f"Folder: {folder} · {len(lines)} entries" + (f" (truncated at {MAX_LIST})" if truncated else "")
        text = head + "\n" + ("\n".join(lines) if lines else "(empty)")
        return FileResult(text, True, "file_listed", str(folder), f"{len(lines)} entries")

    def search_files(self, query: str = "", under: str | None = None) -> FileResult:
        q = (query or "").strip().lower()
        if not q:
            raise FileAccessError("empty query")
        starts = [self.sb.resolve(under)] if (under or "").strip() else [r for r in self.sb.roots if r.is_dir()]
        glob = any(c in q for c in "*?[")
        words = q.split()
        hits: list[str] = []
        seen: set[Path] = set()
        for start in starts:
            for p, st, is_dir, _ in self._walk(start, True):
                if p in seen:
                    continue
                name = p.name.lower()
                if (fnmatch.fnmatch(name, q) if glob else all(w in name for w in words)):
                    seen.add(p)
                    hits.append(f"{p}{'/' if is_dir else ''}\t{'' if is_dir else _size(st.st_size)}\t"
                                f"{_mtime(st.st_mtime)}")
                    if len(hits) >= MAX_HITS:
                        break
            if len(hits) >= MAX_HITS:
                break
        ref = str(starts[0]) if len(starts) == 1 else "(readable folders)"
        head = f"Search '{query}': {len(hits)} match(es)" + (f" (first {MAX_HITS})" if len(hits) >= MAX_HITS else "")
        return FileResult(head + ("\n" + "\n".join(hits) if hits else ""), True, "file_listed", ref,
                          f"search '{_short(query, 60)}': {len(hits)} hits")

    # -- read --

    def read_file(self, path: str = "", offset: Any = 0, max_chars: Any = None, sheet: str | None = None
                  ) -> FileResult:
        p = self.sb.resolve(path)
        if p.is_dir():
            raise FileAccessError(f"'{path}' is a folder (use list_files)")
        text, kind = extract_text(p, sheet or None)
        total = len(text)
        start = max(0, _int(offset, 0))
        limit = self.sb.max_chars
        n = max(1, min(limit, _int(max_chars, limit) if max_chars else limit))
        chunk = text[start:start + n]
        end = start + len(chunk)
        head = f"File: {p} ({kind}, {_size(p.stat().st_size)}) · characters {start:,}–{end:,} of {total:,}"
        if end < total:
            head += f"\n[more text: call read_file with offset={end} to continue]"
        elif start >= total and total:
            head += "\n[offset is past the end of the text]"
        detail = f"chars {start}-{end} of {total}" + (f", sheet {sheet}" if sheet else "")
        return FileResult(head + "\n\n" + chunk, True, "file_read", str(p), detail)

    # -- write --

    def write_deliverable(self, filename: str = "", format: str = "", content: Any = None,
                          sheets: dict[str, Any] | None = None) -> FileResult:
        fmt = str(format or "").lower().lstrip(".") or Path(str(filename)).suffix.lower().lstrip(".")
        data = render_deliverable(fmt, content, sheets)
        requested = _with_ext(paths.safe_name(str(filename or "deliverable")), fmt)
        out = self.sb.outputs
        out.mkdir(parents=True, exist_ok=True)
        stem, ext = requested[: -len(fmt) - 1], "." + fmt
        name, n = requested, 2
        while True:
            target = out / name
            try:
                with open(target, "xb") as fh:  # exclusive create: never overwrite
                    fh.write(data)
                break
            except FileExistsError:
                name = f"{stem} ({n}){ext}"
                n += 1
                if n > 999:
                    raise FileAccessError("too many files with this name") from None
        kind = {"md": "markdown", "json": "json", "txt": "text"}.get(fmt, "file")
        att = Attachment(name=name, kind=kind, size_bytes=len(data),
                         download_url=download_url(self.sb.mission_id, name))
        detail = f"{_size(len(data))}" + (f", requested as {requested}" if name != requested else "")
        text = (f"Wrote {name} ({_size(len(data))}) to the mission outputs. "
                f"Download link for the human: {att.download_url}")
        return FileResult(text, True, "file_written", str(target), detail, attachment=att)


def _with_ext(name: str, fmt: str) -> str:
    if not fmt:
        return name
    stem = name[: -len(Path(name).suffix)] if Path(name).suffix.lower() == "." + fmt else name
    if Path(stem).suffix.lower().lstrip(".") in FORMATS:
        stem = stem[: -len(Path(stem).suffix)]
    return f"{stem or 'deliverable'}.{fmt}"
